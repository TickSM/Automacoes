import tkinter as tk
from tkinter import messagebox, filedialog
import customtkinter as ctk
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from selenium import webdriver
from selenium.webdriver.support.ui import Select, WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from selenium.common.exceptions import NoSuchElementException
import unicodedata
import os
import json
import re
import threading
import time


# ───────────── FUNÇÕES DE APOIO ─────────────
def limpar_nome_arquivo(nome):
    return re.sub(r'[\\/*?:"<>|]', "-", str(nome))


def normalizar_texto(texto):
    if not texto:
        return ""
    texto = str(texto)
    nfkd = unicodedata.normalize('NFKD', texto)
    sem_acento = "".join([c for c in nfkd if not unicodedata.combining(c)])
    sem_acento = re.sub(r'[\u2010-\u2015\u2212]', '-', sem_acento)
    sem_acento = re.sub(r'\s*/\s*', '/', sem_acento)
    return " ".join(sem_acento.replace('\xa0', ' ').lower().split())


def converter_data(data_str):
    """Converte string no formato DD/MM/AAAA para objeto datetime.date"""
    if not data_str:
        return None
    data_limpa = str(data_str).strip()
    match = re.search(r'(\d{2}/\d{2}/\d{4})', data_limpa)
    if match:
        try:
            return datetime.strptime(match.group(1), "%d/%m/%Y").date()
        except Exception:
            pass
    return None


def extrair_data_e_hora(texto):
    """Extrai objeto date e string completa de data/hora a partir de um texto"""
    if not texto:
        return None, None
    texto_limpo = str(texto).strip()
    match = re.search(r'(\d{2}/\d{2}/\d{4})(?:\s+(\d{2}:\d{2}(?::\d{2})?))?', texto_limpo)
    if match:
        dt_obj = converter_data(match.group(1))
        return dt_obj, match.group(0)
    return None, None


def selecionar_opcao(select_element, texto_busca):
    if not texto_busca or not str(texto_busca).strip():
        return
    
    texto_busca_limpo = str(texto_busca).strip()
    select = Select(select_element)
    
    # 1. Tentativa direta com o texto exato
    try:
        select.select_by_visible_text(texto_busca_limpo)
        return
    except Exception:
        pass

    opcoes = select.options
    busca_norm = normalizar_texto(texto_busca_limpo)
    
    # 2. Correspondência exata normalizada (sem acentos, minúsculas, hífens padronizados)
    for opt in opcoes:
        opt_text = opt.text.strip()
        if normalizar_texto(opt_text) == busca_norm:
            val = opt.get_attribute("value")
            if val is not None:
                select.select_by_value(val)
            else:
                select.select_by_visible_text(opt_text)
            return

    # 3. Correspondência parcial (termo procurado contido na opção ou vice-versa)
    candidatos = []
    for opt in opcoes:
        opt_text = opt.text.strip()
        opt_norm = normalizar_texto(opt_text)
        if busca_norm in opt_norm or opt_norm in busca_norm:
            candidatos.append(opt)
    
    if candidatos:
        candidatos.sort(key=lambda o: 0 if normalizar_texto(o.text).startswith(busca_norm) else 1)
        val = candidatos[0].get_attribute("value")
        if val is not None:
            select.select_by_value(val)
        else:
            select.select_by_visible_text(candidatos[0].text.strip())
        return

    # 4. Correspondência por palavras-chave relevantes
    palavras = [p for p in busca_norm.split() if len(p) >= 3 and p != '-']
    if palavras:
        candidatos_palavras = []
        for opt in opcoes:
            opt_norm = normalizar_texto(opt.text)
            if all(p in opt_norm for p in palavras):
                candidatos_palavras.append(opt)
        if candidatos_palavras:
            val = candidatos_palavras[0].get_attribute("value")
            if val is not None:
                select.select_by_value(val)
            else:
                select.select_by_visible_text(candidatos_palavras[0].text.strip())
            return

    # 5. Se nenhuma opção bater, lista opções semelhantes encontradas no menu
    sugestoes = []
    if palavras:
        for opt in opcoes:
            opt_text = opt.text.strip()
            opt_norm = normalizar_texto(opt_text)
            if any(p in opt_norm for p in palavras):
                sugestoes.append(opt_text)

    msg_erro = f"Não foi possível localizar '{texto_busca_limpo}' no menu do SEI."
    if sugestoes:
        msg_erro += "\n\nOpções semelhantes disponíveis:\n" + "\n".join(f"• {s}" for s in sugestoes[:10])
    raise NoSuchElementException(msg_erro)


driver = None


# ───────────── JANELA DE STATUS MODERNA ─────────────
class JanelaStatusModerna(ctk.CTkToplevel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.title("Status da Automação")
        self.geometry("460x220")
        self.resizable(False, False)
        self.attributes("-topmost", True)
        self.protocol("WM_DELETE_WINDOW", self.confirmar_interrupcao)

        # Container principal com cantos arredondados
        self.card = ctk.CTkFrame(self, corner_radius=12, fg_color=("#F0F3F6", "#21252B"))
        self.card.pack(fill="both", expand=True, padx=15, pady=15)

        self.lbl_titulo = ctk.CTkLabel(
            self.card,
            text="⚙️ Automação em Execução",
            font=("Arial", 14, "bold"),
            text_color=("#1F497D", "#5DADE2")
        )
        self.lbl_titulo.pack(pady=(12, 6))

        self.lbl_msg = ctk.CTkLabel(
            self.card,
            text="Iniciando...",
            font=("Arial", 12, "bold"),
            wraplength=400
        )
        self.lbl_msg.pack(pady=(2, 4))

        self.lbl_detalhe = ctk.CTkLabel(
            self.card,
            text="Preparando ambiente e navegador...",
            font=("Arial", 11),
            text_color=("gray40", "gray70"),
            wraplength=400
        )
        self.lbl_detalhe.pack(pady=(0, 10))

        self.progress = ctk.CTkProgressBar(self.card, mode="indeterminate", width=380, height=10)
        self.progress.pack(pady=(5, 12))
        self.progress.start()

    def confirmar_interrupcao(self):
        if messagebox.askyesno("Interromper", "Deseja realmente interromper a automação?", parent=self):
            try:
                if driver:
                    driver.quit()
            except Exception:
                pass
            finally:
                os._exit(0)

    def set_msg(self, texto):
        self.lbl_msg.configure(text=texto)

    def set_detalhe(self, texto):
        self.lbl_detalhe.configure(text=texto)


# Wrapper compatível com o método .config() da interface anterior
class StatusBridge:
    def __init__(self, janela, tipo):
        self.janela = janela
        self.tipo = tipo

    def config(self, text=""):
        if self.tipo == "msg":
            self.janela.set_msg(text)
        elif self.tipo == "detalhe":
            self.janela.set_detalhe(text)


def criar_janela_status(root):
    win = JanelaStatusModerna(root)
    return win, StatusBridge(win, "msg"), StatusBridge(win, "detalhe")


# ───────────── INTERFACE DE CONFIGURAÇÃO MODERNA ─────────────
class AppConfiguracao(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Relatório de Andamento - Verificação de Unidade SEI")
        self.geometry("740x820")
        self.minsize(700, 750)
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        self.CONFIG_FILE = "config_andamento.json"
        self.dados_resultado = None

        config_anterior = self.carregar_config_anterior()

        # Canvas/Scrollable Frame para suportar diferentes resoluções de tela
        self.scroll = ctk.CTkScrollableFrame(self, fg_color="transparent")
        self.scroll.pack(fill="both", expand=True, padx=20, pady=15)

        # Cabeçalho Principal
        self.lbl_titulo = ctk.CTkLabel(
            self.scroll,
            text="📋 Verificação de Andamento no SEI",
            font=("Arial", 20, "bold")
        )
        self.lbl_titulo.pack(pady=(5, 2))

        self.lbl_subtitulo = ctk.CTkLabel(
            self.scroll,
            text="Pesquisa processos, captura a especificação e audita a passagem por unidade e período",
            font=("Arial", 12),
            text_color="gray70"
        )
        self.lbl_subtitulo.pack(pady=(0, 15))

        # CARD 1: Credenciais e Acesso
        self.card_login = self._criar_card("🔐 Acesso ao SEI")
        self.ent_usuario = self._adicionar_input(self.card_login, "Usuário SEI:", config_anterior.get('usuario', ''))
        self.ent_senha = self._adicionar_input(self.card_login, "Senha SEI:", "", is_pass=True)

        # CARD 2: Filtros da Pesquisa do SEI
        self.card_pesquisa = self._criar_card("🔎 Filtros da Pesquisa Inicial")
        self.ent_unidade_geradora = self._adicionar_input(
            self.card_pesquisa, "Unidade Geradora (Opcional):", config_anterior.get('unidade_geradora', '')
        )
        self.ent_tipo_proc = self._adicionar_input(
            self.card_pesquisa, "Tipo de Processo:", config_anterior.get('tipo_processo', '')
        )

        frame_datas_pesq = ctk.CTkFrame(self.card_pesquisa, fg_color="transparent")
        frame_datas_pesq.pack(fill="x", padx=15, pady=4)
        self.ent_dt_ini_pesq = self._adicionar_input_coluna(
            frame_datas_pesq, "Data Início Busca (DD/MM/AAAA):",
            config_anterior.get('data_inicio_pesquisa', config_anterior.get('data_inicio', '')),
            col=0
        )
        self.ent_dt_fim_pesq = self._adicionar_input_coluna(
            frame_datas_pesq, "Data Final Busca (DD/MM/AAAA):",
            config_anterior.get('data_final_pesquisa', config_anterior.get('data_final', '')),
            col=1
        )

        # CARD 3: Auditoria do Andamento
        self.card_andamento = self._criar_card("📍 Verificação no Andamento")
        self.ent_unidade_alvo = self._adicionar_input(
            self.card_andamento, "Unidade Alvo a Verificar:",
            config_anterior.get('unidade_alvo', 'DIAF/GNAF/CCC/Contratos')
        )

        frame_datas_verif = ctk.CTkFrame(self.card_andamento, fg_color="transparent")
        frame_datas_verif.pack(fill="x", padx=15, pady=4)
        self.ent_dt_ini_verif = self._adicionar_input_coluna(
            frame_datas_verif, "Data Início Verificação (DD/MM/AAAA):",
            config_anterior.get('data_inicio_verif', ''),
            col=0
        )
        self.ent_dt_fim_verif = self._adicionar_input_coluna(
            frame_datas_verif, "Data Final Verificação (DD/MM/AAAA):",
            config_anterior.get('data_final_verif', ''),
            col=1
        )

        # CARD 4: Opções e Destino
        self.card_opcoes = self._criar_card("⚙️ Opções & Destino")
        self.chk_processos_var = tk.BooleanVar(value=config_anterior.get('pesquisar_processos', True))
        self.chk_processos = ctk.CTkCheckBox(
            self.card_opcoes,
            text="Pesquisar por Processos (marcar opção 'Processos' no SEI)",
            variable=self.chk_processos_var,
            font=("Arial", 12)
        )
        self.chk_processos.pack(anchor="w", padx=15, pady=(5, 4))

        self.chk_tramitacao_var = tk.BooleanVar(value=config_anterior.get('tramitacao', False))
        self.chk_tramitacao = ctk.CTkCheckBox(
            self.card_opcoes,
            text="Pesquisar com Tramitação na Unidade",
            variable=self.chk_tramitacao_var,
            font=("Arial", 12)
        )
        self.chk_tramitacao.pack(anchor="w", padx=15, pady=(4, 10))

        # Caminho de Salvamento
        lbl_salvar = ctk.CTkLabel(self.card_opcoes, text="Salvar Relatório Excel em:", font=("Arial", 11, "bold"))
        lbl_salvar.pack(anchor="w", padx=15, pady=(2, 2))

        frame_salvar = ctk.CTkFrame(self.card_opcoes, fg_color="transparent")
        frame_salvar.pack(fill="x", padx=15, pady=(0, 10))

        caminho_padrao = config_anterior.get('caminho_salvamento', os.path.join(os.path.expanduser("~"), "Desktop"))
        self.ent_caminho = ctk.CTkEntry(frame_salvar, height=32)
        self.ent_caminho.insert(0, caminho_padrao)
        self.ent_caminho.pack(side="left", fill="x", expand=True, padx=(0, 8))

        self.btn_pasta = ctk.CTkButton(
            frame_salvar,
            text="📁 Procurar",
            width=95,
            height=32,
            fg_color="#34495E",
            hover_color="#2C3E50",
            command=self.selecionar_pasta
        )
        self.btn_pasta.pack(side="right")

        # Botão de Ação Principal
        self.btn_iniciar = ctk.CTkButton(
            self.scroll,
            text="🚀 Iniciar Verificação de Andamento",
            font=("Arial", 14, "bold"),
            height=44,
            fg_color="#1F618D",
            hover_color="#1B4F72",
            corner_radius=8,
            command=self.validar_e_iniciar
        )
        self.btn_iniciar.pack(fill="x", padx=10, pady=(15, 20))

    def _criar_card(self, titulo):
        card = ctk.CTkFrame(self.scroll, corner_radius=10, fg_color=("#EAECEE", "#2B2B2B"))
        card.pack(fill="x", padx=10, pady=7)
        lbl = ctk.CTkLabel(
            card,
            text=titulo,
            font=("Arial", 13, "bold"),
            text_color=("#1F497D", "#5DADE2")
        )
        lbl.pack(anchor="w", padx=15, pady=(10, 6))
        return card

    def _adicionar_input(self, parent, label_text, valor_padrao="", is_pass=False):
        lbl = ctk.CTkLabel(parent, text=label_text, font=("Arial", 11, "bold"))
        lbl.pack(anchor="w", padx=15, pady=(4, 2))
        ent = ctk.CTkEntry(parent, height=32, show="*" if is_pass else "")
        if valor_padrao:
            ent.insert(0, valor_padrao)
        ent.pack(fill="x", padx=15, pady=(0, 6))
        return ent

    def _adicionar_input_coluna(self, parent, label_text, valor_padrao="", col=0):
        frame_col = ctk.CTkFrame(parent, fg_color="transparent")
        frame_col.pack(side="left", fill="x", expand=True, padx=(0 if col == 0 else 6, 6 if col == 0 else 0))
        lbl = ctk.CTkLabel(frame_col, text=label_text, font=("Arial", 11, "bold"))
        lbl.pack(anchor="w", pady=(0, 2))
        ent = ctk.CTkEntry(frame_col, height=32)
        if valor_padrao:
            ent.insert(0, valor_padrao)
        ent.pack(fill="x")
        return ent

    def selecionar_pasta(self):
        pasta = filedialog.askdirectory()
        if pasta:
            self.ent_caminho.delete(0, "end")
            self.ent_caminho.insert(0, pasta)

    def carregar_config_anterior(self):
        try:
            with open(self.CONFIG_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            try:
                with open("config_sei.json", 'r', encoding='utf-8') as f:
                    return json.load(f)
            except Exception:
                return {}

    def validar_e_iniciar(self):
        obrigatorios = {
            "Usuário SEI": self.ent_usuario.get(),
            "Senha SEI": self.ent_senha.get(),
            "Tipo de Processo": self.ent_tipo_proc.get(),
            "Data Início Busca SEI": self.ent_dt_ini_pesq.get(),
            "Data Final Busca SEI": self.ent_dt_fim_pesq.get(),
            "Unidade Alvo no Andamento": self.ent_unidade_alvo.get(),
            "Data Início Verificação": self.ent_dt_ini_verif.get(),
            "Data Final Verificação": self.ent_dt_fim_verif.get()
        }

        vazios = [c for c, v in obrigatorios.items() if not str(v).strip()]
        if vazios:
            messagebox.showwarning(
                "Campos Obrigatórios",
                "Os seguintes campos devem ser preenchidos:\n\n• " + "\n• ".join(vazios),
                parent=self
            )
            return

        datas_validar = [
            ("Data Início Busca SEI", self.ent_dt_ini_pesq.get()),
            ("Data Final Busca SEI", self.ent_dt_fim_pesq.get()),
            ("Data Início Verificação", self.ent_dt_ini_verif.get()),
            ("Data Final Verificação", self.ent_dt_fim_verif.get())
        ]
        for nome_dt, val_dt in datas_validar:
            if converter_data(val_dt) is None:
                messagebox.showwarning(
                    "Data Inválida",
                    f"O campo '{nome_dt}' deve estar no formato DD/MM/AAAA válido.",
                    parent=self
                )
                return

        self.dados_resultado = {
            'usuario': self.ent_usuario.get().strip(),
            'senha': self.ent_senha.get().strip(),
            'unidade_geradora': self.ent_unidade_geradora.get().strip(),
            'tipo_processo': self.ent_tipo_proc.get().strip(),
            'data_inicio_pesquisa': self.ent_dt_ini_pesq.get().strip(),
            'data_final_pesquisa': self.ent_dt_fim_pesq.get().strip(),
            'unidade_alvo': self.ent_unidade_alvo.get().strip(),
            'data_inicio_verif': self.ent_dt_ini_verif.get().strip(),
            'data_final_verif': self.ent_dt_fim_verif.get().strip(),
            'dt_ini_verif': converter_data(self.ent_dt_ini_verif.get()),
            'dt_fim_verif': converter_data(self.ent_dt_fim_verif.get()),
            'pesquisar_processos': self.chk_processos_var.get(),
            'tramitacao': self.chk_tramitacao_var.get(),
            'caminho_salvamento': self.ent_caminho.get().strip()
        }

        # Salva as configurações sem a senha
        config_save = self.dados_resultado.copy()
        del config_save['senha']
        del config_save['dt_ini_verif']
        del config_save['dt_fim_verif']
        try:
            with open(self.CONFIG_FILE, 'w', encoding='utf-8') as f:
                json.dump(config_save, f, indent=4, ensure_ascii=False)
        except Exception:
            pass

        self.destroy()


# ───────────── LÓGICA CORE (SELENIUM) ─────────────
def executar_automacao(dados, lb_msg, lb_detalhe, win_status):
    global driver
    try:
        lb_msg.config(text="Iniciando Navegador...")

        chrome_options = Options()
        # ---- Performance tweaks -------------------------------------------------
        # 1. Headless mode – uncomment the line below if you do not need a visible UI.
        # chrome_options.add_argument("--headless")
        # 2. Faster page load strategy (eager) – loads the DOM without waiting for images.
        chrome_options.page_load_strategy = "eager"
        # 3. Disable images and extensions to reduce bandwidth and rendering time.
        prefs = {
            "profile.managed_default_content_settings.images": 2,
            "profile.default_content_setting_values.notifications": 2,
            "profile.managed_default_content_settings.stylesheets": 2,
        }
        chrome_options.add_experimental_option("prefs", prefs)
        chrome_options.add_argument("--disable-extensions")
        chrome_options.add_argument("--disable-gpu")
        chrome_options.add_argument("--window-size=1920,1080")
        # -----------------------------------------------------------------------

        driver = webdriver.Chrome(options=chrome_options)
        wait = WebDriverWait(driver, 15)

        lb_msg.config(text="Realizando Login no SEI...")
        driver.get("https://sei.assefaz.org.br/sip/login.php?sigla_orgao_sistema=ASSEFAZ&sigla_sistema=SEI")

        driver.find_element(By.ID, "txtUsuario").send_keys(dados['usuario'])
        driver.find_element(By.ID, "pwdSenha").send_keys(dados['senha'])
        driver.find_element(By.ID, "Acessar").click()

        driver.find_element(By.ID, "txtPesquisaRapida").send_keys(Keys.RETURN)

        lb_msg.config(text="Aplicando Filtros de Pesquisa...")

        # Marca opção Processos se selecionado
        if dados.get('pesquisar_processos'):
            try:
                try:
                    label_processos = wait.until(EC.element_to_be_clickable((By.ID, "lblProcessos")))
                    label_processos.click()
                except Exception:
                    radio_processos = wait.until(EC.presence_of_element_located((By.ID, "optProcessos")))
                    if not radio_processos.is_selected():
                        driver.execute_script("arguments[0].click();", radio_processos)
                time.sleep(1)
            except Exception as e:
                print(f"Aviso ao marcar opção Processos: {e}")

        # Filtro de Unidade Geradora
        if dados.get('unidade_geradora'):
            campo_u = driver.find_element(By.ID, "txtUnidade")
            campo_u.send_keys(dados['unidade_geradora'])
            time.sleep(1)
            campo_u.send_keys(Keys.ARROW_DOWN, Keys.ENTER)

        # Filtro de Tipo de Processo
        if dados.get('tipo_processo', '').strip():
            elem_proc = wait.until(EC.presence_of_element_located((By.ID, "selTipoProcedimentoPesquisa")))
            selecionar_opcao(elem_proc, dados['tipo_processo'])

        # Datas da Pesquisa
        driver.find_element(By.ID, "txtDataInicio").send_keys(dados['data_inicio_pesquisa'])
        driver.find_element(By.ID, "txtDataFim").send_keys(dados['data_final_pesquisa'])

        # Tramitação
        if dados.get('tramitacao'):
            try:
                chk_tramitacao = driver.find_element(By.ID, "chkSinTramitacao")
                if not chk_tramitacao.is_selected():
                    driver.execute_script("arguments[0].click();", chk_tramitacao)
            except Exception:
                pass

        driver.find_element(By.ID, 'sbmPesquisar').click()

        processos_detalhes = []
        contador = 0
        unidade_alvo_norm = normalizar_texto(dados['unidade_alvo'])
        dt_ini_verif = dados['dt_ini_verif']
        dt_fim_verif = dados['dt_fim_verif']

        while True:
            try:
                tabela = wait.until(EC.presence_of_element_located((By.CLASS_NAME, "pesquisaResultado")))
                linhas = tabela.find_elements(By.XPATH, './/tr[@class="pesquisaTituloRegistro"]')
            except Exception:
                linhas = []

            if not linhas:
                break

            aba_original = driver.current_window_handle

            for linha in linhas:
                contador += 1
                link = None
                num_proc = "-"
                eh_processo_valido = False

                # Localiza o link do processo pelo NUP
                links_da_linha = linha.find_elements(By.TAG_NAME, "a")
                for l in links_da_linha:
                    href = l.get_attribute("href") or ""
                    texto_link = l.get_attribute("textContent").strip()
                    match_nup = re.search(r'[\w\.]*\d{4,6}/\d{4}-\d{2}', texto_link)
                    if match_nup and "procedimento_trabalhar" in href:
                        link = l
                        num_proc = match_nup.group(0)
                        eh_processo_valido = True
                        break

                if not eh_processo_valido:
                    try:
                        link_generico = linha.find_element(By.XPATH, './/a[contains(@class, "protocoloNormal")]')
                        texto_gen = link_generico.get_attribute("textContent").strip()
                        num_proc = texto_gen if texto_gen else f"Item {contador} (Não identificado)"
                    except Exception:
                        num_proc = f"Item {contador} (Sem link)"

                # Metadados do resultado da pesquisa (Usuário e Data Inclusão)
                try:
                    meta_tr = linha.find_element(By.XPATH, "./following-sibling::tr[2]")
                    usuario_nome = meta_tr.find_element(By.XPATH,
                                                        ".//b[contains(text(), 'Usuário:')]/following-sibling::a").get_attribute(
                        "title")
                    data_inc = meta_tr.find_element(By.XPATH, ".//td[b[contains(text(), 'Inclusão:')]]").get_attribute(
                        "textContent").split(":")[-1].strip().replace('"', '')
                except Exception:
                    usuario_nome = "Não extraído"
                    data_inc = "Não extraído"

                if not eh_processo_valido or not link:
                    processos_detalhes.append([
                        num_proc,
                        "NÃO (Link inválido)",
                        "-",
                        "-",
                        "-",
                        usuario_nome,
                        data_inc
                    ])
                    continue

                lb_msg.config(text=f"Processando ({contador}): {num_proc}")
                lb_detalhe.config(text=f"Abrindo processo {num_proc}...")

                # Abre o processo em nova aba
                driver.execute_script("window.open(arguments[0].href, '_blank');", link)
                driver.switch_to.window(driver.window_handles[-1])

                try:
                    # 1. Extração do campo Especificação (Consultar/Alterar Processo no ifrVisualizacao)
                    especificacao_processo = "-"
                    try:
                        driver.switch_to.default_content()
                        wait.until(EC.frame_to_be_available_and_switch_to_it("ifrVisualizacao"))

                        xpath_consultar_proc = (
                            "//a[contains(., 'Consultar Processo') or contains(@title, 'Consultar Processo') or contains(@aria-label, 'Consultar Processo')] | "
                            "//a[contains(., 'Alterar Processo') or contains(@title, 'Alterar Processo') or contains(@aria-label, 'Alterar Processo')] | "
                            "//a[contains(@href, 'procedimento_consultar') or contains(@href, 'procedimento_alterar')] | "
                            "//a[.//img[contains(@title, 'Consultar Processo') or contains(@title, 'Alterar Processo') or contains(@alt, 'Consultar Processo') or contains(@alt, 'Alterar Processo')]]"
                        )
                        btn_consultar_proc = WebDriverWait(driver, 5).until(EC.presence_of_element_located((By.XPATH, xpath_consultar_proc)))
                        try:
                            btn_consultar_proc.click()
                        except Exception:
                            driver.execute_script("arguments[0].click();", btn_consultar_proc)

                        # Aguarda o campo de especificação carregar
                        campo_espec = WebDriverWait(driver, 5).until(
                            EC.presence_of_element_located((
                                By.XPATH,
                                "//input[@id='txtDescricao'] | //textarea[@id='txtDescricao'] | //input[contains(@name, 'txtDescricao')] | //*[@aria-label='Especificação:' or @title='Especificação:']"
                            ))
                        )
                        especificacao_processo = (campo_espec.get_attribute("value") or campo_espec.text or "").strip()
                        if not especificacao_processo:
                            especificacao_processo = "-"
                    except Exception as e_esp:
                        print(f"[{num_proc}] Aviso ao capturar especificação: {e_esp}")
                        especificacao_processo = "-"

                    # 2. Acessa o iframe da árvore e obtém o botão/link de Consultar Andamento
                    driver.switch_to.default_content()
                    wait.until(EC.frame_to_be_available_and_switch_to_it("ifrArvore"))

                    xpath_andamento = (
                        "//a[contains(@href, 'procedimento_consultar_historico')] | "
                        "//a[contains(@title, 'Consultar Andamento') or contains(@aria-label, 'Consultar Andamento') or contains(., 'Consultar Andamento')] | "
                        "//a[.//img[contains(@title, 'Consultar Andamento') or contains(@alt, 'Consultar Andamento') or contains(@src, 'consultar_andamento')]]"
                    )

                    btn_andamento = wait.until(EC.presence_of_element_located((By.XPATH, xpath_andamento)))
                    href_andamento = btn_andamento.get_attribute("href")

                    # Tenta clicar no botão de andamento
                    try:
                        btn_andamento.click()
                    except Exception:
                        driver.execute_script("arguments[0].click();", btn_andamento)

                    # 3. Acessa o iframe de visualização e garante que a página de histórico foi carregada
                    driver.switch_to.default_content()
                    wait.until(EC.frame_to_be_available_and_switch_to_it("ifrVisualizacao"))

                    # Verifica se o andamento carregou no ifrVisualizacao
                    andamento_carregado = False
                    try:
                        WebDriverWait(driver, 4).until(
                            lambda d: "procedimento_consultar_historico" in d.execute_script("return window.location.href;")
                            or len(d.find_elements(By.XPATH, "//table[contains(@class, 'infraTable')] | //*[contains(text(), 'Atualizar Andamento') or contains(text(), 'histórico completo')]")) > 0
                        )
                        andamento_carregado = True
                    except Exception:
                        andamento_carregado = False

                    # Se o clique não recarregou o iframe com o histórico, força a navegação direta pelo href!
                    if not andamento_carregado and href_andamento:
                        print(f"[{num_proc}] Forçando navegação direta para o andamento via URL...")
                        driver.get(href_andamento)

                    # Aguarda a tabela de histórico carregar completamente
                    WebDriverWait(driver, 12).until(
                        lambda d: len(d.find_elements(By.XPATH, "//table[contains(@class, 'infraTable')] | //tr[td] | //*[contains(text(), 'Atualizar Andamento') or contains(text(), 'histórico completo')]")) > 0
                    )

                    # 3. Clica em 'Ver histórico completo' caso exista
                    try:
                        btn_historico = WebDriverWait(driver, 2.5).until(EC.element_to_be_clickable((
                            By.XPATH,
                            "//a[contains(., 'Ver histórico completo') or contains(@title, 'Ver histórico completo') or contains(text(), 'histórico completo')]"
                        )))
                        try:
                            btn_historico.click()
                        except Exception:
                            driver.execute_script("arguments[0].click();", btn_historico)
                        
                        # Espera inteligente para a tabela atualizar após clicar em ver completo
                        try:
                            WebDriverWait(driver, 3).until(
                                lambda d: len(d.find_elements(By.XPATH, "//table[contains(@class, 'infraTable')]//tr[td]")) > 0
                            )
                        except Exception:
                            time.sleep(0.5)
                    except Exception:
                        pass

                    # 4. Analisa as linhas do histórico para a unidade alvo no período
                    passou_no_periodo = False
                    data_mov_encontrada = "-"
                    unidade_encontrada = "-"
                    data_fora_periodo = None

                    # Script JS que extrai todas as linhas e células em 1 única chamada ultra rápida
                    JS_EXTRAIR_TABELA = """
                        var linhas = document.querySelectorAll("table.infraTable tr, table tr");
                        var resultado = [];
                        for (var i = 0; i < linhas.length; i++) {
                            var tds = linhas[i].querySelectorAll("td");
                            if (tds.length > 0) {
                                var rowTexts = [];
                                for (var j = 0; j < tds.length; j++) {
                                    rowTexts.push(tds[j].textContent || "");
                                }
                                resultado.push({
                                    celulas: rowTexts,
                                    linhaCompleta: linhas[i].textContent || ""
                                });
                            }
                        }
                        return resultado;
                    """

                    while True:
                        # Extrai toda a tabela em uma única viagem de rede
                        dados_linhas = driver.execute_script(JS_EXTRAIR_TABELA)

                        for item in dados_linhas:
                            celulas = item.get("celulas", [])
                            linha_texto = item.get("linhaCompleta", "")
                            
                            unidade_bateu = False
                            texto_unidade_celula = ""

                            for cel_texto in celulas:
                                c_clean = cel_texto.strip()
                                if c_clean and (unidade_alvo_norm in normalizar_texto(c_clean)):
                                    unidade_bateu = True
                                    texto_unidade_celula = c_clean
                                    break

                            # Fallback no texto completo da linha caso a unidade esteja na descrição
                            if not unidade_bateu:
                                if unidade_alvo_norm in normalizar_texto(linha_texto):
                                    unidade_bateu = True
                                    texto_unidade_celula = dados['unidade_alvo']

                            if unidade_bateu:
                                # Procura a data nas células desta linha
                                data_obj = None
                                data_str = None
                                for c_clean in celulas:
                                    dt_o, dt_s = extrair_data_e_hora(c_clean)
                                    if dt_o:
                                        data_obj, data_str = dt_o, dt_s
                                        break

                                if not data_obj:
                                    dt_o, dt_s = extrair_data_e_hora(linha_texto)
                                    if dt_o:
                                        data_obj, data_str = dt_o, dt_s

                                if data_obj:
                                    if dt_ini_verif <= data_obj <= dt_fim_verif:
                                        passou_no_periodo = True
                                        data_mov_encontrada = data_str
                                        unidade_encontrada = texto_unidade_celula
                                        break  # Encontrado dentro do período!
                                    else:
                                        if not data_fora_periodo:
                                            data_fora_periodo = data_str

                        if passou_no_periodo:
                            # Conforme solicitado: se localizar dentro do período, passa para o próximo processo
                            break

                        # Se não achou na página atual, verifica se há paginação de histórico
                        try:
                            btn_prox_hist = driver.find_element(
                                By.XPATH,
                                "//a[contains(text(), 'Próxima') or contains(@title, 'Próxima Página') or contains(@title, 'Próxima')]"
                            )
                            classe_btn = btn_prox_hist.get_attribute("class") or ""
                            if "infraPaginacaoDesabilitada" not in classe_btn:
                                # Guarda texto da primeira linha para saber quando a página virou
                                primeira_linha_velha = dados_linhas[0]["linhaCompleta"] if dados_linhas else ""
                                driver.execute_script("arguments[0].click();", btn_prox_hist)
                                
                                # Aguarda dinamicamente a troca de conteúdo da tabela
                                try:
                                    WebDriverWait(driver, 3).until(
                                        lambda d: (d.execute_script("var r = document.querySelector('table tr td'); return r ? r.textContent : '';") != primeira_linha_velha)
                                    )
                                except Exception:
                                    time.sleep(0.5)
                            else:
                                break
                        except Exception:
                            break

                    if passou_no_periodo:
                        status_str = "SIM"
                    elif data_fora_periodo:
                        status_str = f"NÃO (Passou em {data_fora_periodo})"
                    else:
                        status_str = "NÃO"

                    lb_detalhe.config(text=f"Resultado: {status_str} | Data: {data_mov_encontrada}")

                    processos_detalhes.append([
                        num_proc,
                        status_str,
                        data_mov_encontrada,
                        unidade_encontrada if passou_no_periodo else dados['unidade_alvo'],
                        especificacao_processo,
                        usuario_nome,
                        data_inc
                    ])

                except Exception as e:
                    print(f"Erro ao consultar andamento do processo {num_proc}: {e}")
                    processos_detalhes.append([
                        num_proc,
                        f"ERRO ({type(e).__name__})",
                        "-",
                        dados['unidade_alvo'],
                        especificacao_processo if 'especificacao_processo' in locals() else "-",
                        usuario_nome,
                        data_inc
                    ])
                finally:
                    driver.close()
                    driver.switch_to.window(aba_original)

            # Próxima página da pesquisa de processos
            try:
                btn_p = driver.find_element(By.LINK_TEXT, "Próxima")
                driver.execute_script("arguments[0].click();", btn_p)
                time.sleep(1.5)
            except Exception:
                break

        # ───────────── GERAÇÃO DO EXCEL ─────────────
        lb_msg.config(text="Gerando relatório consolidado em Excel...")
        lb_detalhe.config(text="Formatando dados...")

        wb = Workbook()
        ws = wb.active
        ws.title = "Andamento Unidade"

        # Título
        ws.merge_cells('A1:G1')
        ws['A1'] = "RELATÓRIO DE ANDAMENTO - VERIFICAÇÃO DE PASSAGEM POR UNIDADE"
        ws['A1'].font = Font(bold=True, size=13, color="FFFFFF")
        ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
        ws['A1'].fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        ws.row_dimensions[1].height = 28

        # Metadados da Consulta
        total_analisados = len(processos_detalhes)
        total_sim = sum(1 for p in processos_detalhes if p[1] == "SIM")
        total_nao = total_analisados - total_sim

        meta_info = [
            ("Tipo de Processo:", dados.get('tipo_processo', '-')),
            ("Período de Busca SEI:", f"{dados['data_inicio_pesquisa']} até {dados['data_final_pesquisa']}"),
            ("Unidade Alvo Verificada:", dados['unidade_alvo']),
            ("Período de Verificação no Andamento:", f"{dados['data_inicio_verif']} até {dados['data_final_verif']}"),
            ("Resumo da Auditoria:", f"Total Analisados: {total_analisados} | Passaram na Unidade: {total_sim} | Não Passaram: {total_nao}")
        ]

        for idx, (label, val) in enumerate(meta_info, start=3):
            ws[f'A{idx}'] = label
            ws[f'A{idx}'].font = Font(bold=True, size=10)
            ws[f'B{idx}'] = val
            ws[f'B{idx}'].font = Font(size=10)

        # Cabeçalhos da Tabela
        headers = [
            "NUP SEI",
            "PASSOU NA UNIDADE?",
            "DATA DA MOVIMENTAÇÃO",
            "UNIDADE REGISTRADA",
            "ESPECIFICAÇÃO DO PROCESSO",
            "USUÁRIO GERADOR DO PROCESSO",
            "DATA INCLUSÃO DO PROCESSO"
        ]

        row_header = 9
        ws.row_dimensions[row_header].height = 24
        header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='BFBFBF'),
            right=Side(style='thin', color='BFBFBF'),
            top=Side(style='thin', color='BFBFBF'),
            bottom=Side(style='thin', color='BFBFBF')
        )

        for c_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=row_header, column=c_idx, value=h_text)
            cell.font = Font(bold=True, size=10, color="1F497D")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = header_fill
            cell.border = thin_border

        # Preenchimento dos Dados
        sim_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Verde suave
        nao_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # Laranja/Vermelho suave

        for r_idx, r_data in enumerate(processos_detalhes, start=10):
            ws.row_dimensions[r_idx].height = 20
            for c_idx, val in enumerate(r_data, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.border = thin_border
                cell.font = Font(size=9)
                cell.alignment = Alignment(vertical="center", horizontal="center" if c_idx in [1, 2, 3, 7] else "left")

                # Destaque visual na coluna "PASSOU NA UNIDADE?"
                if c_idx == 2:
                    if str(val) == "SIM":
                        cell.fill = sim_fill
                        cell.font = Font(bold=True, color="276A3C", size=9)
                    else:
                        cell.fill = nao_fill
                        cell.font = Font(color="C00000", size=9)

        # Ajuste da largura das colunas
        col_widths = {1: 28, 2: 24, 3: 22, 4: 35, 5: 35, 6: 35, 7: 22}
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # Salvamento
        unidade_limpa = limpar_nome_arquivo(dados['unidade_alvo'].replace('/', '_'))
        nome_arquivo = f"Relatorio_Andamento_{unidade_limpa}_{datetime.now().strftime('%d-%m-%Y_%H-%M')}.xlsx"
        caminho_final = os.path.join(dados['caminho_salvamento'], nome_arquivo)

        wb.save(caminho_final)

        msg_sucesso = (
            f"Relatório gerado com sucesso!\n\n"
            f"• Total de processos analisados: {total_analisados}\n"
            f"• Passaram pela unidade '{dados['unidade_alvo']}': {total_sim}\n"
            f"• Não passaram no período: {total_nao}\n\n"
            f"Arquivo salvo em:\n{caminho_final}"
        )
        messagebox.showinfo("Sucesso", msg_sucesso)

    except Exception as e:
        import traceback
        erro_completo = traceback.format_exc()
        print(erro_completo)
        messagebox.showerror("Erro na Automação", f"{type(e).__name__}\n\n{erro_completo}")
    finally:
        if driver:
            try:
                driver.quit()
            except Exception:
                pass
        win_status.destroy()


if __name__ == "__main__":
    app = AppConfiguracao()
    app.mainloop()

    if app.dados_resultado:
        config = app.dados_resultado
        win_status, lb_msg, lb_detalhe = criar_janela_status(None)
        threading.Thread(target=executar_automacao, args=(config, lb_msg, lb_detalhe, win_status)).start()
        win_status.mainloop()
