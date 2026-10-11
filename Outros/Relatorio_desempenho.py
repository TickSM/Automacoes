import customtkinter as ctk
import tkinter as tk
import threading
import os
import time
import pandas as pd
from datetime import datetime
import traceback
import sys

from selenium.common.exceptions import NoSuchElementException, TimeoutException, NoAlertPresentException
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from selenium import webdriver
from selenium.common.exceptions import NoSuchElementException, TimeoutException
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

# ---------------------------------------------------------
# CLASSE AUXILIAR: TOOLTIP (BALÃO DE INFORMAÇÃO)
# ---------------------------------------------------------
class ToolTip:
    def __init__(self, widget, text):
        self.widget = widget
        self.text = text
        self.tooltip_window = None
        self.widget.bind("<Enter>", self.show_tooltip)
        self.widget.bind("<Leave>", self.hide_tooltip)

    def show_tooltip(self, event=None):
        x = self.widget.winfo_rootx() + 25
        y = self.widget.winfo_rooty() + 10
        
        self.tooltip_window = tk.Toplevel(self.widget)
        self.tooltip_window.wm_overrideredirect(True)
        self.tooltip_window.wm_geometry(f"+{x}+{y}")
        
        label = tk.Label(self.tooltip_window, text=self.text, justify="left",
                         bg="#2b2b2b", fg="white", relief="solid", bd=1,
                         font=("Arial", 11))
        label.pack(ipadx=8, ipady=4)

    def hide_tooltip(self, event=None):
        if self.tooltip_window:
            self.tooltip_window.destroy()
            self.tooltip_window = None

# ---------------------------------------------------------
# 1. FUNÇÃO DE AUTOMAÇÃO (SELENIUM)
# ---------------------------------------------------------
def executar_relatorio(app_gui, usuario, senha, unidade, data_inicio, data_final, colaborador, caminho_destino):
    driver = None
    try:
        app_gui.atualizar_status("Iniciando o navegador...")
        
        opcoes = Options()
        modo_headless = False 
        
        if modo_headless:
            opcoes.add_argument("--headless=new")
            opcoes.add_argument("--disable-gpu")
            opcoes.add_argument("--no-sandbox")
            opcoes.add_argument("--disable-dev-shm-usage")
            opcoes.add_argument("--window-size=1920,1080")
        else:
            opcoes.add_argument("--start-maximized")
            
        driver = webdriver.Chrome(options=opcoes)
        wait = WebDriverWait(driver, 10)

        app_gui.atualizar_status("Acessando a tela de Login do SEI...")
        driver.get("https://sei.assefaz.org.br/sip/login.php?sigla_orgao_sistema=ASSEFAZ&sigla_sistema=SEI")

        driver.find_element(By.ID, "txtUsuario").send_keys(usuario)
        driver.find_element(By.ID, "pwdSenha").send_keys(senha)
        driver.find_element(By.ID, "Acessar").click()
        time.sleep(2) # Pausa para dar tempo do alerta aparecer (se houver erro)

        # --- VERIFICAÇÃO DE ERRO NO LOGIN ---
        try:
            alerta = driver.switch_to.alert
            texto_erro = alerta.text
            alerta.accept() # Clica em "OK" no alerta para fechar
            raise Exception(f"Não foi possível logar no SEI. Mensagem: {texto_erro}")
        except NoAlertPresentException:
            pass # Se não teve alerta, o login funcionou e o código segue
            
        app_gui.atualizar_status("Login realizado. Aplicando filtros de pesquisa...")

        campo_pesquisa_rapida = driver.find_element(By.ID, "txtPesquisaRapida")
        campo_pesquisa_rapida.send_keys("")
        campo_pesquisa_rapida.send_keys(Keys.RETURN)
        time.sleep(3)

        # --- TRATAMENTO E VALIDAÇÃO DA UNIDADE ---
        campo_unidade_geradora = driver.find_element(By.CSS_SELECTOR, "input#txtUnidade")
        campo_unidade_geradora.click()
        campo_unidade_geradora.send_keys(Keys.CONTROL + "a")
        campo_unidade_geradora.send_keys(Keys.DELETE)
        
        if unidade:
            app_gui.atualizar_status(f"Inserindo a unidade: {unidade}")
            for char in unidade:
                campo_unidade_geradora.send_keys(char)
                time.sleep(0.5)
            campo_unidade_geradora.send_keys(Keys.ARROW_DOWN, Keys.ENTER, Keys.TAB)
            time.sleep(1) # Aguarda o SEI processar a entrada
            
            # Validação: verifica se o campo continuou preenchido
            valor_unidade = campo_unidade_geradora.get_attribute("value").strip()
            if not valor_unidade:
                raise Exception(f"A unidade '{unidade}' não foi encontrada/reconhecida pelo SEI. Busca abortada.")
        else:
            app_gui.atualizar_status("Pesquisa Global (Sem unidade específica).")

        # --- TRATAMENTO E VALIDAÇÃO DO COLABORADOR ---
        try:
            campo_user = driver.find_element(By.ID, "txtUsuarioGerador1")
            campo_user.click()
            campo_user.send_keys(Keys.CONTROL + "a")
            campo_user.send_keys(Keys.DELETE)
            
            if colaborador:
                app_gui.atualizar_status(f"Inserindo o colaborador: {colaborador}")
                for char in colaborador:
                    campo_user.send_keys(char)
                    time.sleep(0.5)
                campo_user.send_keys(Keys.ARROW_DOWN, Keys.ENTER, Keys.TAB)
                time.sleep(1) # Aguarda o SEI processar a entrada
                
                # Validação: verifica se o campo continuou preenchido
                valor_user = campo_user.get_attribute("value").strip()
                if not valor_user:
                    raise Exception(f"O colaborador '{colaborador}' não foi encontrado/reconhecido pelo SEI. Busca abortada.")
            else:
                app_gui.atualizar_status("Pesquisando todos os colaboradores da unidade.")
        except NoSuchElementException:
            raise Exception("Campo de Usuário Gerador não encontrado na tela.")

        # --- DATAS ---
        driver.find_element(By.ID, "txtDataInicio").send_keys(data_inicio)
        driver.find_element(By.ID, "txtDataFim").send_keys(data_final)
        
        app_gui.atualizar_status("Iniciando extração de dados (Isso pode demorar alguns minutos)...")
        botao_pesquisar = driver.find_element(By.ID, 'sbmPesquisar')
        driver.execute_script("arguments[0].click();", botao_pesquisar)
        time.sleep(2)

        processos_detalhes = []
        pagina_atual = 1

        while True:
            try:
                wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "table.pesquisaResultado tbody")))
                linhas_tabela = driver.find_elements(By.CSS_SELECTOR, "table.pesquisaResultado tbody > tr")

                processo_atual = None

                for linha in linhas_tabela:
                    classe_linha = linha.get_attribute("class")

                    if "pesquisaTituloRegistro" in (classe_linha or ""):
                        if processo_atual:
                            processos_detalhes.append(processo_atual)

                        processo_atual = ["Não encontrado", "Não encontrado", "Não especificado", "-", "-", "-"]

                        try:
                            celula_titulo = linha.find_element(By.CSS_SELECTOR, "td.pesquisaTituloEsquerda")
                            processo_atual[0] = celula_titulo.find_element(By.TAG_NAME, "span").text.strip()
                            links_protocolo = celula_titulo.find_elements(By.CSS_SELECTOR, "a.protocoloNormal")
                            if len(links_protocolo) >= 2:
                                processo_atual[1] = links_protocolo[0].text.strip()
                                processo_atual[2] = links_protocolo[1].text.strip()
                            elif len(links_protocolo) == 1:
                                processo_atual[1] = links_protocolo[0].text.strip()
                        except Exception:
                            pass

                    elif processo_atual is not None:
                        try:
                            tds_metatag = linha.find_elements(By.CSS_SELECTOR, "td.pesquisaMetatag")
                            for td in tds_metatag:
                                texto_td = td.text
                                if "Inclusão:" in texto_td:
                                    processo_atual[3] = texto_td.replace("Inclusão:", "").replace('"', '').strip()
                                elif "Unidade:" in texto_td:
                                    processo_atual[4] = texto_td.replace("Unidade:", "").strip()
                                elif "Usuário:" in texto_td:
                                    processo_atual[5] = texto_td.replace("Usuário:", "").strip()
                        except Exception:
                            pass

                if processo_atual:
                    processos_detalhes.append(processo_atual)

                app_gui.atualizar_status(f"Página {pagina_atual} processada. Total acumulado: {len(processos_detalhes)} documentos.")
                
                try:
                    link_proxima = wait.until(EC.element_to_be_clickable((By.LINK_TEXT, "Próxima")))
                    link_proxima.click()
                    pagina_atual += 1
                    time.sleep(2)
                except (NoSuchElementException, TimeoutException):
                    break 

            except Exception as e:
                app_gui.atualizar_status(f"Erro no loop principal de extração: {e}")
                break

        app_gui.atualizar_status(f"\nExtração concluída! Total geral: {len(processos_detalhes)} processos encontrados.")

        if processos_detalhes:
            try:
                os.makedirs(caminho_destino, exist_ok=True)
                
                dados_tabela = []
                if unidade and not colaborador:
                    colunas = ['Nome do Processo', 'Número do Processo', 'Tipo de Documento', 'Data de Inclusão', 'Colaborador']
                    coluna_final = 'E'
                    max_cols = 5
                    for p in processos_detalhes:
                        dados_tabela.append([p[0], p[1], p[2], p[3], p[5]]) 
                elif colaborador and not unidade:
                    colunas = ['Nome do Processo', 'Número do Processo', 'Tipo de Documento', 'Data de Inclusão', 'Unidade']
                    coluna_final = 'E'
                    max_cols = 5
                    for p in processos_detalhes:
                        dados_tabela.append([p[0], p[1], p[2], p[3], p[4]]) 
                else: 
                    colunas = ['Nome do Processo', 'Número do Processo', 'Tipo de Documento', 'Data de Inclusão']
                    coluna_final = 'D'
                    max_cols = 4
                    for p in processos_detalhes:
                        dados_tabela.append([p[0], p[1], p[2], p[3]])

                df = pd.DataFrame(dados_tabela, columns=colunas)
                
                sufixo_arquivo = (unidade.replace('/', '_') if unidade else "Global") + "_" + (colaborador.replace(' ', '_') if colaborador else "Todos")
                data_hora_atual = datetime.now().strftime("%d-%m-%Y_%H-%M-%S")
                nome_arquivo = f"Relatorio_Processos_SEI_{sufixo_arquivo}_{data_hora_atual}.xlsx"
                caminho_completo = os.path.join(caminho_destino, nome_arquivo)

                with pd.ExcelWriter(caminho_completo, engine='openpyxl') as writer:
                    df.to_excel(writer, index=False, startrow=4, sheet_name="Relatório")
                    planilha = writer.sheets["Relatório"]
                    
                    planilha.sheet_view.showGridLines = False

                    planilha.merge_cells(f'A1:{coluna_final}1')
                    a1 = planilha['A1']
                    a1.value = "RELATÓRIO DE DESEMPENHO"
                    a1.font = Font(size=14, bold=True)
                    a1.alignment = Alignment(horizontal='center', vertical='center')

                    try:
                        f_bold = InlineFont(b=True, sz=11)
                        f_normal = InlineFont(b=False, sz=11)
                        
                        planilha['A2'] = CellRichText(TextBlock(f_bold, "Período pesquisado: "), TextBlock(f_normal, f"{data_inicio} a {data_final}"))
                        
                        if unidade and not colaborador:
                            planilha['A3'] = CellRichText(TextBlock(f_bold, "Unidade: "), TextBlock(f_normal, f"{unidade}"))
                            planilha['A4'] = CellRichText(TextBlock(f_bold, "Colaborador: "), TextBlock(f_normal, "Todos"))
                        elif colaborador and not unidade:
                            planilha['A3'] = CellRichText(TextBlock(f_bold, "Colaborador: "), TextBlock(f_normal, f"{colaborador}"))
                            planilha['A4'] = CellRichText(TextBlock(f_bold, "Unidade: "), TextBlock(f_normal, "Todas"))
                        else:
                            planilha['A3'] = CellRichText(TextBlock(f_bold, "Colaborador: "), TextBlock(f_normal, f"{colaborador}"))
                            planilha['A4'] = CellRichText(TextBlock(f_bold, "Unidade: "), TextBlock(f_normal, f"{unidade}"))
                    except Exception:
                        planilha['A2'] = f"Período pesquisado: {data_inicio} a {data_final}"
                        planilha['A3'] = f"Colaborador: {colaborador if colaborador else 'Todos'}"
                        planilha['A4'] = f"Unidade: {unidade if unidade else 'Todas'}"
                        for cell in ['A2', 'A3', 'A4']:
                            planilha[cell].font = Font(bold=True, size=11)

                    planilha.column_dimensions['A'].width = 53.71
                    planilha.column_dimensions['B'].width = 33.29
                    planilha.column_dimensions['C'].width = 47.00
                    planilha.column_dimensions['D'].width = 21.00
                    if max_cols == 5:
                        planilha.column_dimensions['E'].width = 30.00 

                    borda_simples = Border(
                        left=Side(style='thin'), right=Side(style='thin'),
                        top=Side(style='thin'), bottom=Side(style='thin')
                    )
                    alinhamento_central = Alignment(horizontal='center', vertical='center')
                    fonte_titulo = Font(bold=True)
                    fonte_corpo = Font(bold=False)

                    for row in planilha.iter_rows(min_row=5, max_row=planilha.max_row, min_col=1, max_col=max_cols):
                        for cell in row:
                            cell.border = borda_simples
                            cell.alignment = alinhamento_central
                            if cell.row == 5:
                                cell.font = fonte_titulo 
                            else:
                                cell.font = fonte_corpo 

                app_gui.atualizar_status(f"✔ Sucesso! Arquivo Excel gerado em:\n{caminho_completo}")

            except Exception as e:
                app_gui.atualizar_status(f"[ERRO] Não foi possível salvar o arquivo Excel: {e}")
        else:
            app_gui.atualizar_status("A pesquisa não retornou nenhum resultado. O arquivo Excel não foi gerado.")

    except Exception as e:
        app_gui.atualizar_status(f"\n[ERRO] {str(e)}")
        
    finally:
        if driver:
            try:
                driver.quit()
            except:
                pass
        app_gui.finalizar_execucao()

# ---------------------------------------------------------
# 2. INTERFACE GRÁFICA (CustomTkinter)
# ---------------------------------------------------------
class InterfaceRelatorio(ctk.CTk):
    def __init__(self):
        super().__init__()
        
        self.title("Relatório de Desempenho SEI")
        self.geometry("580x630") 
        ctk.set_appearance_mode("dark")
        
        self.lbl_titulo = ctk.CTkLabel(self, text="Gerador de Relatório de Desempenho", font=("Arial", 18, "bold"))
        self.lbl_titulo.pack(pady=(15, 10))
        
        self.frame_inputs = ctk.CTkFrame(self, fg_color="transparent")
        self.frame_inputs.pack(pady=5)
        
        ctk.CTkLabel(self.frame_inputs, text="Usuário SEI:").grid(row=0, column=0, padx=5, pady=5, sticky="e")
        self.input_usuario = ctk.CTkEntry(self.frame_inputs, width=220)
        self.input_usuario.grid(row=0, column=1, padx=5, pady=5)
        lbl_info_usuario = ctk.CTkLabel(self.frame_inputs, text="(?)", text_color="#aaaaaa", cursor="hand2")
        lbl_info_usuario.grid(row=0, column=2, padx=5, sticky="w")
        ToolTip(lbl_info_usuario, "Seu login utilizado para acessar o sistema SEI.")
        
        ctk.CTkLabel(self.frame_inputs, text="Senha:").grid(row=1, column=0, padx=5, pady=5, sticky="e")
        self.input_senha = ctk.CTkEntry(self.frame_inputs, width=220, show="*")
        self.input_senha.grid(row=1, column=1, padx=5, pady=5)
        lbl_info_senha = ctk.CTkLabel(self.frame_inputs, text="(?)", text_color="#aaaaaa", cursor="hand2")
        lbl_info_senha.grid(row=1, column=2, padx=5, sticky="w")
        ToolTip(lbl_info_senha, "Sua senha de acesso ao sistema SEI.")

        ctk.CTkLabel(self.frame_inputs, text="Data Inicial (dd/mm/aaaa):").grid(row=2, column=0, padx=5, pady=5, sticky="e")
        self.input_data_inicio = ctk.CTkEntry(self.frame_inputs, width=220)
        self.input_data_inicio.grid(row=2, column=1, padx=5, pady=5)
        lbl_info_inicio = ctk.CTkLabel(self.frame_inputs, text="(?)", text_color="#aaaaaa", cursor="hand2")
        lbl_info_inicio.grid(row=2, column=2, padx=5, sticky="w")
        ToolTip(lbl_info_inicio, "Obrigatório. Data de início da pesquisa. Ex: 01/01/2024")
        
        ctk.CTkLabel(self.frame_inputs, text="Data Final (dd/mm/aaaa):").grid(row=3, column=0, padx=5, pady=5, sticky="e")
        self.input_data_final = ctk.CTkEntry(self.frame_inputs, width=220)
        self.input_data_final.grid(row=3, column=1, padx=5, pady=5)
        lbl_info_fim = ctk.CTkLabel(self.frame_inputs, text="(?)", text_color="#aaaaaa", cursor="hand2")
        lbl_info_fim.grid(row=3, column=2, padx=5, sticky="w")
        ToolTip(lbl_info_fim, "Obrigatório. Data de término da pesquisa. Ex: 31/12/2024")

        ctk.CTkLabel(self.frame_inputs, text="Unidade:").grid(row=4, column=0, padx=5, pady=5, sticky="e")
        self.input_unidade = ctk.CTkEntry(self.frame_inputs, width=220)
        self.input_unidade.grid(row=4, column=1, padx=5, pady=5)
        lbl_info_unidade = ctk.CTkLabel(self.frame_inputs, text="(?)", text_color="#aaaaaa", cursor="hand2")
        lbl_info_unidade.grid(row=4, column=2, padx=5, sticky="w")
        ToolTip(lbl_info_unidade, "Opcional. Sigla exata da unidade geradora. Ex: CDI/ADM\n(Deixe em branco para pesquisar todas as unidades)")

        ctk.CTkLabel(self.frame_inputs, text="Colaborador (Usuário Gerador):").grid(row=5, column=0, padx=5, pady=5, sticky="e")
        self.input_colaborador = ctk.CTkEntry(self.frame_inputs, width=220)
        self.input_colaborador.grid(row=5, column=1, padx=5, pady=5)
        lbl_info_colaborador = ctk.CTkLabel(self.frame_inputs, text="(?)", text_color="#aaaaaa", cursor="hand2")
        lbl_info_colaborador.grid(row=5, column=2, padx=5, sticky="w")
        ToolTip(lbl_info_colaborador, "Opcional. Nome exato do usuário.\n(Deixe em branco para pesquisar todos os colaboradores)")

        ctk.CTkLabel(self.frame_inputs, text="Pasta de Destino (Excel):").grid(row=6, column=0, padx=5, pady=5, sticky="e")
        self.input_pasta = ctk.CTkEntry(self.frame_inputs, width=220)
        self.input_pasta.grid(row=6, column=1, padx=5, pady=5)
        
        lbl_info_pasta = ctk.CTkLabel(self.frame_inputs, text="(?)", text_color="#aaaaaa", cursor="hand2")
        lbl_info_pasta.grid(row=6, column=2, padx=5, sticky="w")
        ToolTip(lbl_info_pasta, "Local no seu computador ou rede onde\no arquivo Excel final será salvo.")
        
        self.btn_pasta = ctk.CTkButton(self.frame_inputs, text="📁", width=30, command=self.selecionar_pasta)
        self.btn_pasta.grid(row=6, column=3, padx=5, pady=5)
        
        self.caixa_logs = ctk.CTkTextbox(self, width=510, height=150)
        self.caixa_logs.pack(pady=10)
        self.caixa_logs.insert("0.0", "Preencha a Unidade OU o Colaborador (ou ambos). Em seguida, clique em Iniciar.\n\n")
        
        self.btn_iniciar = ctk.CTkButton(self, text="Iniciar Extração", command=self.iniciar_thread)
        self.btn_iniciar.pack(pady=5)

    def selecionar_pasta(self):
        pasta_selecionada = ctk.filedialog.askdirectory(title="Selecione a pasta para salvar o Excel")
        if pasta_selecionada:
            self.input_pasta.delete(0, "end")
            self.input_pasta.insert(0, pasta_selecionada)

    def iniciar_thread(self):
        usuario = self.input_usuario.get().strip()
        senha = self.input_senha.get().strip()
        data_inicio = self.input_data_inicio.get().strip()
        data_final = self.input_data_final.get().strip()
        unidade = self.input_unidade.get().strip()
        colaborador = self.input_colaborador.get().strip()
        caminho_destino = self.input_pasta.get().strip()
        
        # VALIDAÇÕES
        if not usuario or not senha or not data_inicio or not data_final or not caminho_destino:
            self.atualizar_status("ERRO: Preencha Usuário, Senha, Datas e Destino antes de iniciar!")
            return
            
        if not unidade and not colaborador:
            self.atualizar_status("ERRO: Você deve preencher pelo menos a Unidade ou o Colaborador!")
            return
            
        self.btn_iniciar.configure(state="disabled") 
        self.atualizar_status("Preparando execução em segundo plano...")
        
        thread_robo = threading.Thread(target=executar_relatorio, args=(
            self, usuario, senha, unidade, data_inicio, data_final, colaborador, caminho_destino
        ))
        thread_robo.start()

    def atualizar_status(self, mensagem):
        self.after(0, self._inserir_texto, mensagem)

    def _inserir_texto(self, mensagem):
        self.caixa_logs.insert("end", mensagem + "\n")
        self.caixa_logs.see("end")

    def finalizar_execucao(self):
        self.after(0, lambda: self.btn_iniciar.configure(state="normal", text="Executar Novamente"))

if __name__ == "__main__":
    app = InterfaceRelatorio()
    app.mainloop()