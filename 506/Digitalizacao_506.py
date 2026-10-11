import os
import sys
import re
import time
import json
import threading
from datetime import datetime

import tkinter as tk
from tkinter import messagebox, ttk, filedialog
import customtkinter as ctk

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import Select, WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ───────────── FUNÇÕES DE APOIO ─────────────
def limpar_nome_arquivo(nome):
    return re.sub(r'[\\/*?:"<>|]', "-", nome)

def buscar_caminho_padrao_planilha():
    """Verifica se a pasta existe nos diretórios corporativos J: ou I:."""
    caminho_empresa = r"J:\Atual\ARQUIVO_2026\DIGITALIZAÇÃO 506\Acompanhamento"
    caminho_casa = r"I:\Atual\ARQUIVO_2026\DIGITALIZAÇÃO 506\Acompanhamento"

    if os.path.exists(caminho_empresa):
        return caminho_empresa
    elif os.path.exists(caminho_casa):
        return caminho_casa
    return ""

driver = None

# ───────────── INTERFACE DO MONITOR (LOG) ─────────────
def criar_janela_status(root, stop_callback):
    status_win = tk.Toplevel(root)
    status_win.title("Monitor da Automação")
    status_win.geometry("650x450")
    status_win.attributes("-topmost", True)

    # Barra de progresso para indicar atividade
    progress = ttk.Progressbar(status_win, mode='indeterminate')
    progress.pack(fill='x', padx=20, pady=(10, 5))
    progress.start()

    # Caixa de Log
    log_frame = ttk.Frame(status_win)
    log_frame.pack(fill='both', expand=True, padx=20, pady=5)

    log_text = tk.Text(log_frame, state='disabled', wrap='word', font=("Consolas", 9))
    scrollbar = ttk.Scrollbar(log_frame, command=log_text.yview)
    log_text.configure(yscrollcommand=scrollbar.set)

    log_text.pack(side='left', fill='both', expand=True)
    scrollbar.pack(side='right', fill='y')

    # Função que injeta mensagens no Log
    def append_log(msg):
        log_text.config(state='normal')
        agora = datetime.now().strftime("%H:%M:%S")
        log_text.insert('end', f"[{agora}] {msg}\n")
        log_text.see('end')
        log_text.config(state='disabled')
        status_win.update_idletasks()

    def handle_stop():
        btn_stop.config(state='disabled', text="Salvando e Encerrando...")
        append_log("⚠️ Sinal de interrupção enviado! O robô vai parar e salvar os dados em breve...")
        stop_callback()

    # Botão de Parada
    btn_stop = tk.Button(
        status_win, 
        text="🛑 PARAR AUTOMAÇÃO E SALVAR", 
        bg="#c0392b", 
        fg="white", 
        font=("Arial", 11, "bold"), 
        command=handle_stop
    )
    btn_stop.pack(pady=10)

    def on_close():
        if messagebox.askyesno("Interromper", "Deseja parar a automação e salvar a planilha?", parent=status_win):
            handle_stop()

    status_win.protocol("WM_DELETE_WINDOW", on_close)

    return status_win, append_log, progress, btn_stop


# ───────────── LÓGICA CORE (SELENIUM) ─────────────
def executar_automacao(dados, append_log, stop_event, win_status, progress, btn_stop):
    global driver
    try:
        append_log("Iniciando verificação de segurança...")

        nome_f = limpar_nome_arquivo(dados['nome_formulario'])
        
        # Define o caminho
        if dados['caminho_salvamento'].lower().endswith(('.xlsx', '.xls')):
            caminho_final = dados['caminho_salvamento']
        else:
            caminho_final = os.path.join(dados['caminho_salvamento'], f"Relatorio_Consolidado_{nome_f}.xlsx")

        nups_ja_extraidos = set()

        # ==========================================
        # TRAVA DE SEGURANÇA DA PLANILHA
        # ==========================================
        if os.path.exists(caminho_final):
            append_log(f"Planilha base localizada: {caminho_final}")
            try:
                wb_existente = load_workbook(caminho_final, read_only=True)
                ws_existente = wb_existente.worksheets[0]

                for row in ws_existente.iter_rows(min_row=9, min_col=1, max_col=1, values_only=True):
                    if row[0]:
                        nups_ja_extraidos.add(str(row[0]).strip())

                wb_existente.close()
                append_log(f"[{len(nups_ja_extraidos)}] Processos encontrados na planilha e serão ignorados.")
            except Exception as e:
                append_log(f"Aviso: Não foi possível ler o arquivo existente. Erro: {e}")
        else:
            append_log(">>> ERRO: A PLANILHA NÃO FOI ENCONTRADA! <<<")
            append_log(f"O arquivo esperado não existe: {caminho_final}")
            append_log("TRAVA DE SEGURANÇA ATIVADA: A automação requer uma planilha existente para continuar.")
            messagebox.showerror(
                "Planilha Ausente", 
                f"Planilha não localizada em:\n{caminho_final}\n\nAutomação cancelada.", 
                parent=win_status
            )
            return  # ABORTA A EXECUÇÃO AQUI!

        append_log("Iniciando Navegador...")
        chrome_options = Options()
        chrome_options.add_argument("--window-size=1920,1080")
        chrome_options.add_argument("--disable-gpu")

        driver = webdriver.Chrome(options=chrome_options)
        wait = WebDriverWait(driver, 15)

        append_log("Realizando Login no SEI...")
        driver.get("https://sei.assefaz.org.br/sip/login.php?sigla_orgao_sistema=ASSEFAZ&sigla_sistema=SEI")

        driver.find_element(By.ID, "txtUsuario").send_keys(dados['usuario'])
        driver.find_element(By.ID, "pwdSenha").send_keys(dados['senha'])
        driver.find_element(By.ID, "Acessar").click()

        append_log("Verificando unidade atual...")
        try:
            driver.switch_to.default_content()
            unidade_atual_elem = wait.until(EC.presence_of_element_located((By.ID, "lnkInfraUnidade")))
            unidade_atual_texto = ""
            for _ in range(6):
                txt_content = unidade_atual_elem.get_attribute("textContent") or ""
                txt_inner = unidade_atual_elem.get_attribute("innerText") or ""
                txt_title = unidade_atual_elem.get_attribute("title") or ""
                txt_selenium = unidade_atual_elem.text or ""

                texto_combinado = f"{txt_content} {txt_inner} {txt_title} {txt_selenium}".strip()
                if texto_combinado:
                    unidade_atual_texto = texto_combinado
                    break
                time.sleep(0.5)

            if "ADB" not in unidade_atual_texto.upper():
                append_log("  ↳ Unidade incorreta. Trocando para ADB...")
                driver.execute_script("arguments[0].click();", unidade_atual_elem)
                time.sleep(2)

                try:
                    label_unidade_adb = wait.until(EC.presence_of_element_located(
                        (By.XPATH, "//tr[contains(., 'ADB') and contains(., 'Arquivo Digital')]//label")
                    ))
                    driver.execute_script("arguments[0].click();", label_unidade_adb)
                except Exception:
                    link_unidade_adb = wait.until(EC.presence_of_element_located(
                        (By.XPATH, "//*[contains(text(), 'ADB') and contains(text(), 'Arquivo Digital')]")
                    ))
                    driver.execute_script("arguments[0].click();", link_unidade_adb)

                time.sleep(3)
                append_log("✔ Unidade alterada para ADB com sucesso!")
            else:
                append_log("✔ Usuário já está na unidade ADB!")

        except Exception as e:
            append_log(f"Erro ao trocar de unidade: {e}")

        driver.find_element(By.ID, "txtPesquisaRapida").send_keys(Keys.RETURN)

        append_log("Aplicando Filtros de Pesquisa...")
        if dados['unidade_geradora']:
            campo_u = driver.find_element(By.ID, "txtUnidade")
            campo_u.send_keys(dados['unidade_geradora'])
            time.sleep(1)
            campo_u.send_keys(Keys.ARROW_DOWN, Keys.ENTER)

        opt_processos = driver.find_element(By.XPATH, "//label[@for='optProcessos']")
        driver.execute_script("arguments[0].click();", opt_processos)
        Select(driver.find_element(By.ID, "selTipoProcedimentoPesquisa")).select_by_visible_text(dados['tipo_processo'])
        driver.find_element(By.ID, "txtDataInicio").send_keys(dados['data_inicio'])
        driver.find_element(By.ID, "txtDataFim").send_keys(dados['data_final'])

        if dados.get('tramitacao'):
            try:
                chk_tramitacao = driver.find_element(By.ID, "chkSinTramitacao")
                if not chk_tramitacao.is_selected():
                    driver.execute_script("arguments[0].click();", chk_tramitacao)
            except Exception:
                pass

        driver.find_element(By.ID, 'sbmPesquisar').click()

        processos_detalhes = []
        contador = 0
        extrair_tudo = "TODOS" in [c.upper() for c in dados['lista_campos_excel']]

        # --- FUNÇÃO DE SALVAMENTO PARCIAL (CHECKPOINT) ---
        def salvar_parcial(forcar_atualizacao=False):
            if not processos_detalhes and not forcar_atualizacao:
                return

            data_hora_atual = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
            try:
                if os.path.exists(caminho_final):
                    wb = load_workbook(caminho_final)
                    ws = wb.worksheets[0]
                    ws['B7'] = data_hora_atual
                    linha_inicio = ws.max_row + 1
                    
                    if processos_detalhes:
                        for r_idx, r_data in enumerate(processos_detalhes, linha_inicio):
                            for c_idx, val in enumerate(r_data, 1):
                                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                                cell.alignment = Alignment(wrap_text=True, vertical="top")

                        for col in ws.columns:
                            ws.column_dimensions[get_column_letter(col[0].column)].width = 40

                    wb.save(caminho_final)
                    qnt_salva = len(processos_detalhes)
                    processos_detalhes.clear()  
                    append_log(f"[CHECKPOINT] {qnt_salva} registros recém-extraídos salvos com sucesso!")
                else:
                    append_log("[ERRO GRAVE] O arquivo Excel base foi excluído ou movido durante a execução.")
            except Exception as e_save:
                append_log(f"Erro ao tentar salvar a planilha parcial: {e_save}")
        # -------------------------------------------------

        while True:
            if stop_event.is_set():
                append_log("⚠️️ Loop de paginação interrompido pelo usuário.")
                break

            tabela = wait.until(EC.presence_of_element_located((By.CLASS_NAME, "pesquisaResultado")))
            linhas = tabela.find_elements(By.XPATH, './/tr[@class="pesquisaTituloRegistro"]')
            aba_original = driver.current_window_handle

            for linha in linhas:
                if stop_event.is_set():
                    break

                contador += 1
                link = None
                num_proc = "-"
                eh_processo_valido = False

                links_da_linha = linha.find_elements(By.TAG_NAME, "a")

                for l in links_da_linha:
                    href = l.get_attribute("href") or ""
                    texto_link = l.get_attribute("textContent").strip()
                    match_nup = re.search(r'[\w\.]*\d{4,6}/\d{4}-\d{2}', texto_link)

                    if match_nup and "procedimento_trabalhar" in href:
                        link = l
                        num_proc = match_nup.group(0)
                        eh_processo_valido = True
                        break

                if not eh_processo_valido:
                    continue

                try:
                    meta_tr = linha.find_element(By.XPATH, "./following-sibling::tr[2]")
                    usuario_nome = meta_tr.find_element(By.XPATH, ".//b[contains(text(), 'Usuário:')]/following-sibling::a").get_attribute("title")
                    data_inc = meta_tr.find_element(By.XPATH, ".//td[b[contains(text(), 'Inclusão:')]]").get_attribute("textContent").split(":")[-1].strip().replace('"', '')
                except Exception:
                    usuario_nome = "Não extraído"
                    data_inc = "Não extraído"

                if num_proc.strip() in nups_ja_extraidos:
                    append_log(f"Pulando {num_proc} (Já presente na planilha)")
                    continue

                append_log(f"Processando: {num_proc}")

                driver.execute_script("window.open(arguments[0].href, '_blank');", link)
                driver.switch_to.window(driver.window_handles[-1])

                try:
                    wait.until(EC.frame_to_be_available_and_switch_to_it("ifrArvore"))

                    while True:
                        pastas = driver.find_elements(By.XPATH, "//img[starts-with(@title, 'Abrir Pasta')]")
                        if not pastas:
                            break
                        for p in pastas:
                            try:
                                driver.execute_script("arguments[0].click();", p)
                                try:
                                    WebDriverWait(driver, 10).until(EC.invisibility_of_element_located((By.XPATH, "//span[contains(@id, 'spanAGUARDE')]")))
                                except Exception:
                                    pass
                            except Exception:
                                pass
                        time.sleep(0.3)

                    xpath_geral = "//a[contains(@class, 'infraArvoreNo') and not(contains(@title, 'Documento Cancelado'))]"
                    todos_docs = driver.find_elements(By.XPATH, xpath_geral)

                    links_f = []
                    for doc in todos_docs:
                        try:
                            texto_span = doc.find_element(By.XPATH, ".//span").get_attribute("textContent").strip()
                            if "Acervo Arq de Dossiê de Beneficiário ADB" in texto_span and "Inativo" in texto_span:
                                links_f.append({"elemento": doc, "status": "INATIVO"})
                            elif "Acervo Arquivístico de Dossiê de Beneficiário ADB" in texto_span:
                                links_f.append({"elemento": doc, "status": "ATIVO"})
                        except Exception:
                            pass

                    if links_f:
                        ids_vistos = []
                        f_count = 0
                        for item in links_f:
                            if stop_event.is_set():
                                break 

                            l_form = item["elemento"]
                            status_form = item["status"]

                            id_l = l_form.get_attribute("id")
                            if id_l in ids_vistos:
                                continue
                            ids_vistos.append(id_l)

                            dados_f = {}
                            f_count += 1
                            append_log(f"  ↳ Lendo documento {f_count}/{len(links_f)} ({status_form})")

                            try:
                                driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", l_form)
                                driver.execute_script("arguments[0].click();", l_form)

                                driver.switch_to.default_content()
                                wait.until(EC.frame_to_be_available_and_switch_to_it("ifrVisualizacao"))
                                wait.until(EC.frame_to_be_available_and_switch_to_it("ifrArvoreHtml"))
                                time.sleep(0.5)

                                elemento_conteudo = wait.until(EC.presence_of_element_located((By.ID, "conteudo")))
                                texto_bruto = elemento_conteudo.get_attribute("innerText")
                                linhas_txt = [l.strip() for l in texto_bruto.splitlines()]

                                if extrair_tudo and (not dados['lista_campos_excel'] or dados['lista_campos_excel'] == ["TODOS"]):
                                    campos_detectados = []
                                    for l in linhas_txt:
                                        if l.endswith(":") and len(l) <= 45 and not l.isupper():
                                            campos_detectados.append(l)
                                    if campos_detectados:
                                        dados['lista_campos_excel'] = campos_detectados

                                for i, txt in enumerate(linhas_txt):
                                    if txt in dados['lista_campos_excel']:
                                        valores = []
                                        j = i + 1
                                        while j < len(linhas_txt):
                                            ln = linhas_txt[j]
                                            if ln in dados['lista_campos_excel'] or (ln.endswith(":") and len(ln) <= 45 and not ln.isupper() and not ln.startswith(("1 -", "2 -", "3 -", "4 -"))):
                                                break
                                            valores.append(ln)
                                            j += 1
                                        dados_f[txt] = "\n".join(valores).strip() if valores else "-"

                                linha_final = [num_proc, usuario_nome, data_inc, status_form]
                                for campo in dados['lista_campos_excel']:
                                    linha_final.append(dados_f.get(campo, "-"))
                                processos_detalhes.append(linha_final)

                            except Exception as e:
                                append_log(f"  ↳ Erro na leitura do documento: {type(e).__name__}")
                                linha_erro = [num_proc, usuario_nome, data_inc, "ERRO"]
                                for _ in dados['lista_campos_excel']:
                                    linha_erro.append(f"ERRO NA EXTRAÇÃO ({type(e).__name__})")
                                processos_detalhes.append(linha_erro)

                            finally:
                                driver.switch_to.default_content()
                                try:
                                    wait.until(EC.frame_to_be_available_and_switch_to_it("ifrArvore"))
                                except Exception:
                                    pass
                    else:
                        append_log(f"  ↳ Nenhum formulário ADB válido encontrado.")

                finally:
                    driver.close()
                    driver.switch_to.window(aba_original)

            salvar_parcial()

            if stop_event.is_set():
                break

            try:
                btn_p = driver.find_element(By.LINK_TEXT, "Próxima")
                driver.execute_script("arguments[0].click();", btn_p)
                time.sleep(1.5)
            except Exception:
                append_log("Fim dos resultados da pesquisa.")
                break

        if stop_event.is_set():
            append_log("Gerando salvamento final devido à interrupção...")
        else:
            append_log("Atualizando relatório excel pela última vez...")

        salvar_parcial(forcar_atualizacao=True)
        append_log("✔ Processamento finalizado! Os dados estão salvos e seguros.")

    except Exception as e:
        append_log("⚠️ ERRO CRÍTICO! Tentando salvar dados residuais...")
        try:
            salvar_parcial()
        except Exception:
            pass

        import traceback
        erro_completo = traceback.format_exc()
        append_log(f"Falha de sistema: {type(e).__name__}")
        messagebox.showerror("Erro", f"{type(e).__name__}\n\n{erro_completo}", parent=win_status)
        
    finally:
        # Fechar o navegador de forma segura
        if driver:
            try:
                driver.quit()
            except Exception:
                pass
        
        # Parar a barra de carregamento visual
        progress.stop()

        # Função para fechar completamente o aplicativo
        def fechar_app():
            os._exit(0)

        # Transforma o botão "Parar" em um botão "Fechar"
        btn_stop.config(
            text="FECHAR APLICATIVO",
            bg="#27ae60", # Cor Verde indicando sucesso e segurança para sair
            state="normal",
            command=fechar_app
        )
        
        # Garante que se o usuário clicar no 'X' da janela, ela também apenas feche
        win_status.protocol("WM_DELETE_WINDOW", fechar_app)


# ───────────── INTERFACE GRÁFICA INICIAL (CustomTkinter) ─────────────
class InterfaceInicialSEI(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Automação Digitalização 506 - SEI")
        self.geometry("500x420")
        self.resizable(False, False)
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        self.lbl_titulo = ctk.CTkLabel(self, text="Digitalização 506 - SEI", font=("Arial", 20, "bold"))
        self.lbl_titulo.pack(pady=(25, 5))

        self.lbl_subtitulo = ctk.CTkLabel(self, text="Informe os dados para iniciar o processamento", font=("Arial", 12), text_color="gray70")
        self.lbl_subtitulo.pack(pady=(0, 15))

        self.frame_form = ctk.CTkFrame(self, fg_color="transparent")
        self.frame_form.pack(pady=5, padx=20, fill="x")

        ctk.CTkLabel(self.frame_form, text="Usuário SEI:", font=("Arial", 13, "bold")).grid(row=0, column=0, padx=10, pady=10, sticky="e")
        self.ent_usuario = ctk.CTkEntry(self.frame_form, width=270, placeholder_text="nome.sobrenome")
        self.ent_usuario.grid(row=0, column=1, padx=5, pady=10)

        ctk.CTkLabel(self.frame_form, text="Senha:", font=("Arial", 13, "bold")).grid(row=1, column=0, padx=10, pady=10, sticky="e")
        self.ent_senha = ctk.CTkEntry(self.frame_form, width=270, show="*")
        self.ent_senha.grid(row=1, column=1, padx=5, pady=10)

        ctk.CTkLabel(self.frame_form, text="Pasta / Planilha:", font=("Arial", 13, "bold")).grid(row=2, column=0, padx=10, pady=10, sticky="e")
        
        self.frame_caminho = ctk.CTkFrame(self.frame_form, fg_color="transparent")
        self.frame_caminho.grid(row=2, column=1, padx=5, pady=10, sticky="w")

        self.ent_caminho = ctk.CTkEntry(self.frame_caminho, width=225)
        self.ent_caminho.pack(side="left", padx=(0, 5))

        self.btn_browse = ctk.CTkButton(self.frame_caminho, text="📁", width=40, command=self.selecionar_caminho)
        self.btn_browse.pack(side="left")

        caminho_padrao = buscar_caminho_padrao_planilha()
        if caminho_padrao:
            self.ent_caminho.insert(0, caminho_padrao)        

        self.btn_iniciar = ctk.CTkButton(
            self, 
            text="Iniciar Automação", 
            font=("Arial", 14, "bold"),
            height=42,
            width=260,
            command=self.iniciar_processamento
        )
        self.btn_iniciar.pack(pady=5)

    def selecionar_caminho(self):
        caminho = filedialog.askdirectory(title="Selecione a pasta de Acompanhamento")
        if not caminho:
            caminho = filedialog.askopenfilename(
                title="Ou selecione a planilha Excel",
                filetypes=[("Arquivos Excel", "*.xlsx *.xls")]
            )
        if caminho:
            self.ent_caminho.delete(0, "end")
            self.ent_caminho.insert(0, caminho)

    def iniciar_processamento(self):
        usuario = self.ent_usuario.get().strip()
        senha = self.ent_senha.get().strip()
        caminho_salvamento = self.ent_caminho.get().strip()

        if not usuario or not senha:
            messagebox.showwarning("Atenção", "Por favor, preencha o Usuário e a Senha!")
            return

        if not caminho_salvamento or not os.path.exists(caminho_salvamento):
            messagebox.showwarning("Atenção", "O caminho da pasta/planilha informado não é válido ou não foi encontrado.")
            return

        config = {
            'usuario': usuario,
            'senha': senha,
            'unidade_geradora': '',
            'tipo_processo': 'GRC - Acervo Arquivístico de Dossiê de Beneficiário',
            'tipo_documento': 'Acervo Arquivístico de Dossiê de Beneficiário ADB',
            'nome_formulario': 'Acervo Arquivístico de Dossiê de Beneficiário ADB',
            'data_inicio': '01/07/2026',
            'data_final': '31/12/2029',
            'tramitacao': False,
            'caminho_salvamento': caminho_salvamento,
            'lista_campos_excel': [
                'MÓDULO:', 'LADO:', 'PRATELEIRA:', 'CAIXA :', 
                'DATA DA INATIVAÇÃO/CANCELAMENTO:', 'NOME DO BENEFICIÁRIO:', 'CPF:'
            ]
        }

        self.withdraw()

        self.stop_event = threading.Event()

        def stop_callback():
            self.stop_event.set()

        win_status, append_log, progress, btn_stop = criar_janela_status(self, stop_callback)

        thread_robo = threading.Thread(
            target=executar_automacao, 
            args=(config, append_log, self.stop_event, win_status, progress, btn_stop),
            daemon=True
        )
        thread_robo.start()

if __name__ == "__main__":
    app = InterfaceInicialSEI()
    app.mainloop()