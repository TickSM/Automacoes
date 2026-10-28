import os
import re
import time
import threading
import traceback
from datetime import datetime

import tkinter as tk
from tkinter import messagebox, ttk, filedialog
import customtkinter as ctk

from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from selenium import webdriver
from selenium.common.exceptions import NoSuchElementException, TimeoutException
from selenium.webdriver.support.ui import Select, WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

driver = None

# ───────────── FUNÇÕES DE APOIO ─────────────
def buscar_caminho_planilha():
    """Busca a planilha nos caminhos padrão. Retorna o caminho se existir, senão None."""
    caminho_j = r"J:\Atual\ARQUIVO_2026\DIGITALIZAÇÃO GR\Acompanhamento\AFD_processos_GR.xlsx"
    caminho_i = r"I:\Atual\ARQUIVO_2026\DIGITALIZAÇÃO GR\Acompanhamento\AFD_processos_GR.xlsx"
    
    if os.path.exists(caminho_j):
        return caminho_j
    elif os.path.exists(caminho_i):
        return caminho_i
    return None

# ───────────── INTERFACE DO MONITOR (LOG) ─────────────
def criar_janela_status(root, stop_callback):
    status_win = tk.Toplevel(root)
    status_win.title("Monitor da Automação - GR")
    status_win.geometry("680x480")
    status_win.attributes("-topmost", True)

    progress = ttk.Progressbar(status_win, mode='indeterminate')
    progress.pack(fill='x', padx=20, pady=(10, 5))
    progress.start()

    log_frame = ttk.Frame(status_win)
    log_frame.pack(fill='both', expand=True, padx=20, pady=5)

    log_text = tk.Text(log_frame, state='disabled', wrap='word', font=("Consolas", 9))
    scrollbar = ttk.Scrollbar(log_frame, command=log_text.yview)
    log_text.configure(yscrollcommand=scrollbar.set)

    log_text.pack(side='left', fill='both', expand=True)
    scrollbar.pack(side='right', fill='y')

    def append_log(msg):
        log_text.config(state='normal')
        agora = datetime.now().strftime("%H:%M:%S")
        log_text.insert('end', f"[{agora}] {msg}\n")
        log_text.see('end')
        log_text.config(state='disabled')
        status_win.update_idletasks()

    def handle_stop():
        btn_stop.config(state='disabled', text="Salvando e Encerrando...")
        append_log("⚠️ Sinal de interrupção enviado! O robô vai finalizar com segurança e salvar os dados...")
        stop_callback()

    btn_stop = tk.Button(
        status_win, 
        text="🛑 PARAR AUTOMAÇÃO E SALVAR", 
        bg="#c0392b", 
        fg="white", 
        font=("Arial", 11, "bold"), 
        command=handle_stop
    )
    btn_stop.pack(pady=10)

    def on_close():
        if messagebox.askyesno("Interromper", "Deseja parar a automação e salvar a planilha?", parent=status_win):
            handle_stop()

    status_win.protocol("WM_DELETE_WINDOW", on_close)

    return status_win, append_log, progress, btn_stop

# ───────────── LÓGICA PRINCIPAL (SELENIUM) ─────────────
def executar_automacao_gr(config, append_log, stop_event, win_status, progress, btn_stop):
    global driver
    try:
        append_log("Iniciando verificação de segurança...")
        
        caminho_excel = config['caminho_excel']
        nups_ja_extraidos = set()

        if not os.path.exists(caminho_excel):
            append_log(">>> ERRO: A PLANILHA NÃO FOI ENCONTRADA! <<<")
            messagebox.showerror("Erro", "O arquivo da planilha desapareceu. Automação cancelada.", parent=win_status)
            return

        # Lê processos existentes para não duplicar
        try:
            wb_existente = load_workbook(caminho_excel, read_only=True)
            ws_existente = wb_existente.worksheets[0]
            for row in ws_existente.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
                if row[0]:
                    nups_ja_extraidos.add(str(row[0]).strip())
            wb_existente.close()
            append_log(f"Encontrados {len(nups_ja_extraidos)} processos já extraídos na planilha.")
        except Exception as e:
            append_log(f"Aviso: Não foi possível ler a planilha existente. Erro: {e}")

        append_log("Abrindo navegador e acessando o SEI...")
        opcoes = Options()
        opcoes.add_argument("--window-size=1920,1080")
        opcoes.add_argument("--disable-gpu")
        
        driver = webdriver.Chrome(options=opcoes)
        wait = WebDriverWait(driver, 10)

        driver.get("https://sei.assefaz.org.br/sip/login.php?sigla_orgao_sistema=ASSEFAZ&sigla_sistema=SEI")

        append_log("Realizando login...")
        driver.find_element(By.ID, "txtUsuario").send_keys(config['usuario'])
        driver.find_element(By.ID, "pwdSenha").send_keys(config['senha'])
        driver.find_element(By.ID, "Acessar").click()
        time.sleep(3)

        append_log("Verificando unidade logada...")
        unidade_ja_logada = False
        try:
            driver.find_element(By.XPATH, "//a[@id='lnkInfraUnidade' and contains(text(), 'DIPRE/GRC/CDI/ADM')]")
            unidade_ja_logada = True
        except NoSuchElementException:
            pass

        if not unidade_ja_logada:
            unidade = driver.find_element(By.XPATH, "(//a[@id='lnkInfraUnidade'])[2]")
            unidade.click()
            time.sleep(1)
            radio = driver.find_element(By.ID, "chkInfraItem143")
            driver.execute_script("arguments[0].click();", radio)
            time.sleep(1.5)

        processos_detalhes = []
        processos_com_erro = []

        # Função de Salvamento Parcial
        def salvar_parcial(forcar_atualizacao=False):
            if not processos_detalhes and not forcar_atualizacao:
                return

            try:
                wb = load_workbook(caminho_excel)
                ws = wb.worksheets[0]
                linha_inicio = ws.max_row + 1

                for r_idx, (nome, nup, und, usuario, data_inc, status) in enumerate(processos_detalhes, start=linha_inicio):
                    ws.cell(row=r_idx, column=1, value=nome)
                    ws.cell(row=r_idx, column=2, value=nup)
                    ws.cell(row=r_idx, column=3, value=und)
                    ws.cell(row=r_idx, column=4, value=usuario)
                    ws.cell(row=r_idx, column=5, value=data_inc)
                    ws.cell(row=r_idx, column=6, value=status)

                for col in ws.columns:
                    max_length = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        try:
                            if cell.value and len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except Exception:
                            pass
                    ws.column_dimensions[col_letter].width = (max_length + 2)

                wb.save(caminho_excel)
                append_log(f"[CHECKPOINT] {len(processos_detalhes)} registros salvos com sucesso!")
                processos_detalhes.clear()
            except Exception as e_save:
                append_log(f"Erro ao tentar salvar a planilha parcial: {e_save}")

        # Lista de UFs a processar
        lista_ufs = config.get('lista_ufs', ['RJ'])
        total_ufs = len(lista_ufs)
        append_log(f"UFs na fila de execução: {', '.join(lista_ufs)} (Total: {total_ufs})")

        estatisticas_execucao = {}

        # ───────────── LOOP PRINCIPAL: ITERAÇÃO POR TODAS AS UFs ─────────────
        for idx_uf, sigla_uf in enumerate(lista_ufs, 1):
            if stop_event.is_set():
                append_log(f"⚠️ Processamento interrompido. UFs pendentes a partir de {sigla_uf} não serão executadas.")
                break

            sigla_formatada = sigla_uf.strip().upper()
            unidade_completa = f"AFD GR/{sigla_formatada}"
            append_log("\n" + "=" * 55)
            append_log(f"📍 [{idx_uf}/{total_ufs}] Regional: {sigla_formatada} ({unidade_completa})")
            append_log("=" * 55)

            novos_extraidos_uf = 0

            # 1. Navega para a tela de pesquisa do SEI
            try:
                # Aceita eventuais alertas pendentes
                try:
                    alert = driver.switch_to.alert
                    alert.accept()
                except Exception:
                    pass

                driver.switch_to.default_content()

                campo_pesquisa_rapida = wait.until(EC.presence_of_element_located((By.ID, "txtPesquisaRapida")))
                campo_pesquisa_rapida.clear()
                campo_pesquisa_rapida.send_keys(Keys.RETURN)
                time.sleep(2.5)
            except Exception as e_nav:
                append_log(f"Aviso na navegação: {e_nav}. Recarregando página de pesquisa...")
                driver.get("https://sei.assefaz.org.br/sei/controlador.php?acao=pesquisa_processo_pesquisar")
                time.sleep(2)

            # 2. Aplica os filtros de pesquisa para a UF atual
            try:
                # Usa o ID da label e o clique nativo do Selenium
                label_processos = wait.until(EC.element_to_be_clickable((By.ID, "lblProcessos")))
                label_processos.click()

                campo_unidade_geradora = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, "input#txtUnidade")))
                campo_unidade_geradora.clear()
                campo_unidade_geradora.click()
                for char in unidade_completa:
                    campo_unidade_geradora.send_keys(char)
                    time.sleep(0.2)

                time.sleep(1.5)
                campo_unidade_geradora.send_keys(Keys.ARROW_DOWN, Keys.ENTER, Keys.TAB)

                elemento_select = wait.until(EC.presence_of_element_located((By.ID, "selTipoProcedimentoPesquisa")))
                dropdown = Select(elemento_select)
                dropdown.select_by_value("100000470")

                campo_data_ini = driver.find_element(By.ID, "txtDataInicio")
                campo_data_ini.clear()
                campo_data_ini.send_keys(config['data_inicio'])

                campo_data_fim = driver.find_element(By.ID, "txtDataFim")
                campo_data_fim.clear()
                campo_data_fim.send_keys(config['data_final'])

                driver.find_element(By.ID, 'sbmPesquisar').click()
                time.sleep(2)
            except Exception as e_filtros:
                append_log(f"❌ Erro ao aplicar filtros para {sigla_formatada}: {e_filtros}")
                estatisticas_execucao[sigla_formatada] = "Erro nos filtros"
                continue

            # Verifica se apareceu alert do SEI (ex: 'Nenhum processo encontrado')
            try:
                alert = driver.switch_to.alert
                msg_alerta = alert.text
                alert.accept()
                append_log(f"Mensagem do SEI para {sigla_formatada}: {msg_alerta}")
                estatisticas_execucao[sigla_formatada] = 0
                continue
            except Exception:
                pass

            # 3. Paginação de resultados da UF atual
            pagina_atual = 1
            while True:
                if stop_event.is_set():
                    append_log("⚠️ Paginação interrompida pelo usuário.")
                    break

                try:
                    conteudo = driver.find_element(By.ID, "conteudo")
                    elementos_processos = conteudo.find_elements(By.XPATH, './/a[@class="protocoloNormal"]')

                    if not elementos_processos:
                        append_log(f"Nenhum processo listado na página {pagina_atual} para {sigla_formatada}.")
                        break

                    aba_original = driver.current_window_handle

                    for elemento in elementos_processos:
                        if stop_event.is_set():
                            break

                        numero_processo = elemento.text.strip()

                        if numero_processo in nups_ja_extraidos:
                            append_log(f"Pulando {numero_processo} - Já extraído.")
                            continue

                        append_log(f"Processando: {numero_processo}")

                        try:
                            unidade_elem = elemento.find_element(By.XPATH, "../../following-sibling::tr//b[contains(text(), 'Unidade:')]/following-sibling::a")
                            unidade_texto = unidade_elem.text.strip()
                        except NoSuchElementException:
                            unidade_texto = "-"

                        try:
                            td_inclusao = elemento.find_element(By.XPATH, "../../following-sibling::tr//td[b[contains(text(), 'Inclusão:')]]")
                            texto_inclusao = td_inclusao.get_attribute("textContent")
                            data_inclusao = texto_inclusao.split(":")[-1].strip().replace('"', '')[:10]
                        except NoSuchElementException:
                            data_inclusao = "-"

                        try:
                            usuario_elem = elemento.find_element(By.XPATH, "../../following-sibling::tr//b[contains(text(), 'Usuário:')]/following-sibling::a")
                            usuario_nome = usuario_elem.get_attribute("title").strip()
                            if not usuario_nome:
                                usuario_nome = usuario_elem.text.strip()
                        except NoSuchElementException:
                            usuario_nome = "-"

                        try:
                            driver.execute_script("arguments[0].click();", elemento)
                        except Exception:
                            processos_com_erro.append(numero_processo)
                            continue

                        wait.until(lambda d: len(d.window_handles) > 1)
                        abas = driver.window_handles
                        driver.switch_to.window(abas[-1])

                        try:
                            driver.switch_to.frame("ifrVisualizacao")
                            formulario = wait.until(EC.element_to_be_clickable((By.XPATH, "//*[@alt='Consultar Processo']")))
                            formulario.click()

                            nome = wait.until(EC.presence_of_element_located((By.XPATH, "//*[@id='txtDescricao']")))
                            valor_nome = nome.get_attribute("value")

                            try:
                                xpath_observacao = "//th[text()='Observação']/parent::tr/following-sibling::tr[1]/td[2]"
                                elem_status = driver.find_element(By.XPATH, xpath_observacao)
                                texto_bruto = elem_status.get_attribute("innerText").strip().upper()
                                
                                if "2 PROJETO DIGITALIZAÇÃO" in texto_bruto or "2 PROJETO DE DIGITALIZAÇÃO" in texto_bruto:
                                    valor_status = "INATIVO" if "INATIVO" in texto_bruto else "ATIVO"
                                    append_log(f"  ↳ [INCLUÍDO] Pertence à 2ª etapa do projeto ({valor_status}).")
                                else:
                                    append_log(f"  ↳ [IGNORADO] Não pertence à 2ª etapa do projeto de digitalização.")
                                    continue 
                            except NoSuchElementException:
                                append_log(f"  ↳ [IGNORADO] Campo de observação vazio ou inexistente.")
                                continue
                                
                            processos_detalhes.append([valor_nome, numero_processo, unidade_texto, usuario_nome, data_inclusao, valor_status])
                            nups_ja_extraidos.add(numero_processo)
                            novos_extraidos_uf += 1

                        except Exception as e:
                            append_log(f"  ↳ Erro no processo {numero_processo}: {type(e).__name__}")
                            processos_com_erro.append(numero_processo)
                        finally:
                            driver.switch_to.default_content()
                            driver.close()
                            driver.switch_to.window(aba_original)

                    # Salva o progresso ao final de cada página
                    salvar_parcial()
                    
                    if stop_event.is_set():
                        break

                    try:
                        link_proxima = wait.until(EC.element_to_be_clickable((By.LINK_TEXT, "Próxima")))
                        driver.execute_script("arguments[0].click();", link_proxima)
                        pagina_atual += 1
                        time.sleep(1.5)
                    except (NoSuchElementException, TimeoutException):
                        break

                except Exception as e:
                    append_log(f"Erro na navegação da tabela: {e}")
                    break

            estatisticas_execucao[sigla_formatada] = novos_extraidos_uf
            append_log(f"✔ Finalizada UF {sigla_formatada}. Novos registros salvos: {novos_extraidos_uf}")
            salvar_parcial(forcar_atualizacao=True)

        # ───────────── RESUMO FINAL DE TODAS AS UFs ─────────────
        if stop_event.is_set():
            append_log("Gerando salvamento final devido à interrupção...")
        else:
            append_log("Atualizando relatório final da execução...")

        salvar_parcial(forcar_atualizacao=True)

        append_log("\n" + "=" * 55)
        append_log("🏁 RELATÓRIO CONSOLIDADO DO PROCESSAMENTO:")
        for uf_item, qtd_item in estatisticas_execucao.items():
            append_log(f"   • {uf_item}: {qtd_item} processo(s) novo(s)")
        append_log("=" * 55)

        if processos_com_erro:
            append_log(f"Atenção: {len(processos_com_erro)} processos apresentaram erro.")
            pasta_destino = os.path.dirname(caminho_excel)
            with open(os.path.join(pasta_destino, "erros_extracao.txt"), "w") as f:
                f.write("\n".join(processos_com_erro))

        append_log("✔ Processamento finalizado! Os dados estão salvos e seguros na planilha.")

    except Exception as e:
        append_log("⚠️ ERRO CRÍTICO! Tentando salvar dados residuais...")
        try:
            salvar_parcial()
        except Exception:
            pass
        append_log(f"Falha de sistema: {type(e).__name__}")
        messagebox.showerror("Erro", traceback.format_exc(), parent=win_status)
    finally:
        if driver:
            try:
                driver.quit()
            except Exception:
                pass
        
        progress.stop()

        def fechar_app():
            os._exit(0)

        btn_stop.config(
            text="FECHAR APLICATIVO",
            bg="#27ae60",
            state="normal",
            command=fechar_app
        )
        win_status.protocol("WM_DELETE_WINDOW", fechar_app)


# ───────────── INTERFACE GRÁFICA (CustomTkinter) ─────────────
class InterfaceInicialGR(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Automação Digitalização GR - SEI")
        self.geometry("490x510")
        self.resizable(False, False)
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        self.lbl_titulo = ctk.CTkLabel(self, text="Digitalização GR - SEI", font=("Arial", 20, "bold"))
        self.lbl_titulo.pack(pady=(20, 5))

        self.lbl_subtitulo = ctk.CTkLabel(
            self, 
            text="Informe suas credenciais e selecione as Regionais (UFs)", 
            font=("Arial", 12), 
            text_color="gray70"
        )
        self.lbl_subtitulo.pack(pady=(0, 15))

        self.frame_form = ctk.CTkFrame(self, fg_color="transparent")
        self.frame_form.pack(pady=5, padx=25, fill="x")

        # Campo Usuário
        ctk.CTkLabel(self.frame_form, text="Usuário SEI:", font=("Arial", 13, "bold")).grid(row=0, column=0, padx=10, pady=8, sticky="e")
        self.ent_usuario = ctk.CTkEntry(self.frame_form, width=230, placeholder_text="nome.sobrenome")
        self.ent_usuario.grid(row=0, column=1, padx=5, pady=8)

        # Campo Senha
        ctk.CTkLabel(self.frame_form, text="Senha:", font=("Arial", 13, "bold")).grid(row=1, column=0, padx=10, pady=8, sticky="e")
        self.ent_senha = ctk.CTkEntry(self.frame_form, width=230, show="*")
        self.ent_senha.grid(row=1, column=1, padx=5, pady=8)

        # Divisor / Título da Seção de UFs
        ctk.CTkLabel(
            self.frame_form, 
            text="Regionais (UFs):", 
            font=("Arial", 13, "bold")
        ).grid(row=2, column=0, padx=10, pady=(15, 5), sticky="ne")

        # Frame de Seleção de UFs
        self.frame_ufs_container = ctk.CTkFrame(self.frame_form, fg_color="#2b2b2b", corner_radius=8)
        self.frame_ufs_container.grid(row=2, column=1, padx=5, pady=(15, 5), sticky="w")

        # Checkboxes das UFs padrão (RJ, SP, BA) pré-marcadas
        self.chk_vars = {}
        frame_chk = ctk.CTkFrame(self.frame_ufs_container, fg_color="transparent")
        frame_chk.pack(padx=10, pady=(8, 4), fill="x")

        ufs_padrao = ["RJ", "SP", "BA"]
        for uf in ufs_padrao:
            var = ctk.BooleanVar(value=True)
            chk = ctk.CTkCheckBox(frame_chk, text=uf, variable=var, font=("Arial", 12, "bold"), width=60)
            chk.pack(side="left", padx=4)
            self.chk_vars[uf] = var

        # Botões de atalho: Todas / Nenhuma
        frame_atalhos = ctk.CTkFrame(self.frame_ufs_container, fg_color="transparent")
        frame_atalhos.pack(padx=10, pady=(2, 8), fill="x")

        btn_todas = ctk.CTkButton(
            frame_atalhos, 
            text="Marcar Todas", 
            width=90, 
            height=24, 
            font=("Arial", 10), 
            command=self.marcar_todas
        )
        btn_todas.pack(side="left", padx=(0, 5))

        btn_nenhuma = ctk.CTkButton(
            frame_atalhos, 
            text="Desmarcar", 
            width=80, 
            height=24, 
            font=("Arial", 10), 
            fg_color="#555555", 
            hover_color="#444444", 
            command=self.desmarcar_todas
        )
        btn_nenhuma.pack(side="left")

        # Campo para adicionar novas UFs no futuro
        ctk.CTkLabel(
            self.frame_form, 
            text="Outras UFs:", 
            font=("Arial", 13, "bold")
        ).grid(row=3, column=0, padx=10, pady=10, sticky="e")

        self.ent_outras_ufs = ctk.CTkEntry(
            self.frame_form, 
            width=230, 
            placeholder_text="Ex: MG, PR, DF (separadas por vírgula)"
        )
        self.ent_outras_ufs.grid(row=3, column=1, padx=5, pady=10)

        # Informação amigável
        self.lbl_info = ctk.CTkLabel(
            self, 
            text="💡 O robô processará todas as UFs selecionadas em sequência\nautomaticamente, com apenas um clique.", 
            font=("Arial", 11), 
            text_color="#85c1e9"
        )
        self.lbl_info.pack(pady=(10, 5))

        self.btn_iniciar = ctk.CTkButton(
            self, 
            text="🚀 Iniciar Automação", 
            font=("Arial", 14, "bold"),
            height=42,
            width=280,
            command=self.iniciar_processamento
        )
        self.btn_iniciar.pack(pady=15)

    def marcar_todas(self):
        for var in self.chk_vars.values():
            var.set(True)

    def desmarcar_todas(self):
        for var in self.chk_vars.values():
            var.set(False)

    def obter_ufs_selecionadas(self):
        ufs = [uf for uf, var in self.chk_vars.items() if var.get()]
        
        texto_outras = self.ent_outras_ufs.get().strip()
        if texto_outras:
            partes = re.split(r'[,;\s]+', texto_outras)
            for parte in partes:
                sigla = parte.strip().upper()
                if sigla and sigla not in ufs:
                    ufs.append(sigla)
        return ufs

    def iniciar_processamento(self):
        usuario = self.ent_usuario.get().strip()
        senha = self.ent_senha.get().strip()
        lista_ufs = self.obter_ufs_selecionadas()

        if not usuario or not senha:
            messagebox.showwarning("Atenção", "Por favor, preencha o Usuário e a Senha!")
            return

        if not lista_ufs:
            messagebox.showwarning("Atenção", "Por favor, selecione ao menos uma UF (ou digite no campo Outras UFs)!")
            return

        # Verifica e valida o caminho da planilha
        caminho_excel = buscar_caminho_planilha()
        
        if not caminho_excel:
            messagebox.showinfo(
                "Planilha não encontrada", 
                "A planilha padrão 'AFD_processos_GR.xlsx' não foi encontrada nas pastas J: ou I:.\n\n"
                "Por favor, localize e selecione o arquivo manualmente na próxima tela."
            )
            caminho_excel = filedialog.askopenfilename(
                title="Selecione a planilha AFD_processos_GR.xlsx",
                filetypes=[("Arquivos Excel", "*.xlsx *.xls")]
            )
            
            if not caminho_excel:
                messagebox.showwarning("Cancelado", "A automação requer a seleção de uma planilha para continuar.")
                return

        config = {
            'usuario': usuario,
            'senha': senha,
            'data_inicio': '27/07/2026', # Data travada
            'data_final': '31/12/2029',  # Data travada
            'caminho_excel': caminho_excel,
            'lista_ufs': lista_ufs
        }

        self.withdraw()
        self.stop_event = threading.Event()

        def stop_callback():
            self.stop_event.set()

        win_status, append_log, progress, btn_stop = criar_janela_status(self, stop_callback)

        thread_robo = threading.Thread(
            target=executar_automacao_gr, 
            args=(config, append_log, self.stop_event, win_status, progress, btn_stop),
            daemon=True
        )
        thread_robo.start()

if __name__ == "__main__":
    app = InterfaceInicialGR()
    app.mainloop()